"""
batch3 — restructure the Waterfall inputs and add a Waterfall Summary tab (modeled on the
55 Resort reference).

The monthly waterfall engine (Waterfall rows 20-50) and the partner IRR/MOIC block (rows
14-18) are unchanged — they reference the assumptions by name — so this only:
  * rebuilds the top of the Waterfall tab into two clean input tables:
      Contributions   : LP / GP  with Amount and %  (GP co-invest % is the single input)
      Distributions   : per-tier  IRR Hurdle | % to GP | % to LP
    and repoints the 7 named assumptions to the new cells, and
  * adds a "Waterfall Summary" tab: a Project/LP/GP contribution-distribution summary
    (Contributions, Distributions, Net Proceeds, IRR, MOIC), a tier table (LP, LP%, GP, GP%,
    Total at tier), and annual LP / GP / GP-incl-fees net proceeds.

Value-neutral: the engine math and the named assumption values are unchanged.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

NAVY = "FF1F3864"; INK = "FF000000"; BLUE = "FF0000CC"; GREEN = "FF006100"; GREY = "FF595959"
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_IN = PatternFill("solid", fgColor="FFFDF3D0")
FILL_TOT = PatternFill("solid", fgColor="FFE2EFDA")
F_TITLE = Font(name="Calibri", size=15, bold=True, color=NAVY)
F_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
F_LBL = Font(name="Calibri", size=10, color=INK)
F_IN = Font(name="Calibri", size=10, color=BLUE)
F_LINK = Font(name="Calibri", size=10, color=GREEN)
F_TOT = Font(name="Calibri", size=10, bold=True, color=INK)
F_HINT = Font(name="Calibri", size=9, italic=True, color=GREY)
RIGHT = Alignment(horizontal="right"); LEFT = Alignment(horizontal="left"); CENTER = Alignment(horizontal="center")
TOPB = Border(top=Side(style="thin", color="FF808080"))
USD = '$#,##0;($#,##0);"–"'; PCT = '0.0%'; PCT2 = '0.00%'; MULT = '0.00"x"'
FIRST, LAST = 5, 376
# Waterfall engine rows
R = dict(tier1=32, tier2=33, tier3=34, tier4=35, gpdist=38, lpdist=39, gpcon=40, lpcon=41,
         gpnet=42, lpnet=43, gpfee=48)


def build(wb):
    wf = wb['Waterfall']

    def repoint(name, ref):
        if name in wb.defined_names: del wb.defined_names[name]
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    def S(ws, coord, v, font=None, fill=None, num=None, align=None, border=None):
        c = ws[coord]; c.value = v
        if font: c.font = font
        if fill: c.fill = fill
        if num: c.number_format = num
        if align: c.alignment = align
        if border: c.border = border
        return c

    # ---------------------------------------------------------- restructure Waterfall inputs
    for r in range(4, 13):
        for col in 'BCDEFGHI':
            wf[f'{col}{r}'].value = None; wf[f'{col}{r}'].fill = PatternFill()
    wf.column_dimensions['F'].width = 30
    for col in ('G', 'H', 'I'):
        wf.column_dimensions[col].width = 12

    # Contributions table (B4:D8)
    S(wf, 'B4', 'Contributions', F_HDR, FILL_HDR, align=LEFT)
    S(wf, 'C4', 'Amount', F_HDR, FILL_HDR, align=RIGHT); S(wf, 'D4', '%', F_HDR, FILL_HDR, align=RIGHT)
    S(wf, 'B5', 'LP', F_LBL, align=LEFT)
    S(wf, 'C5', '=EquityTotal*(1-GPCoInvest)', F_LINK, num=USD, align=RIGHT)
    S(wf, 'D5', '=1-GPCoInvest', F_LINK, num=PCT, align=RIGHT)
    S(wf, 'B6', 'GP', F_LBL, align=LEFT)
    S(wf, 'C6', '=EquityTotal*GPCoInvest', F_LINK, num=USD, align=RIGHT)
    S(wf, 'D6', 0.10, F_IN, FILL_IN, PCT, RIGHT)                 # GP co-invest % (input)
    repoint('GPCoInvest', "'Waterfall'!$D$6")
    S(wf, 'B7', 'Total Contributions', F_TOT, None, None, LEFT, TOPB)
    S(wf, 'C7', '=EquityTotal', F_TOT, None, USD, RIGHT, TOPB)
    S(wf, 'D7', 1, F_TOT, None, PCT, RIGHT, TOPB)
    S(wf, 'B8', 'Check: LP + GP − Total', F_HINT, align=LEFT)
    S(wf, 'C8', '=C5+C6-EquityTotal', F_HINT, num=USD, align=RIGHT)
    S(wf, 'D8', '=IF(ABS(C8)<1,"TRUE","FALSE")', F_HINT, align=RIGHT)

    # Distributions tier table (F4:I8)
    S(wf, 'F4', 'Distributions', F_HDR, FILL_HDR, align=LEFT)
    S(wf, 'G4', 'IRR Hurdle', F_HDR, FILL_HDR, align=RIGHT)
    S(wf, 'H4', '% to GP', F_HDR, FILL_HDR, align=RIGHT)
    S(wf, 'I4', '% to LP', F_HDR, FILL_HDR, align=RIGHT)
    tiers = [
        ('Tier 1 — Return of Capital + Pref', 0.08, 'PrefReturn', '=GPCoInvest', None),
        ('Tier 2 — Pref to Hurdle 2',         0.12, 'Hurdle2IRR', 0.20, 'Promote2'),
        ('Tier 3 — Hurdle 2 to Hurdle 3',     0.15, 'Hurdle3IRR', 0.30, 'Promote3'),
        ('Tier 4 — Residual (> Hurdle 3)',    None, None,         0.40, 'Promote4'),
    ]
    r = 5
    for lab, hv, hname, gpv, gpname in tiers:
        S(wf, f'F{r}', lab, F_LBL, align=LEFT)
        if hname:
            S(wf, f'G{r}', hv, F_IN, FILL_IN, PCT2, RIGHT); repoint(hname, f"'Waterfall'!$G${r}")
        else:
            S(wf, f'G{r}', '="> "&TEXT(Hurdle3IRR,"0.0%")', F_HINT, align=RIGHT)
        if gpname:
            S(wf, f'H{r}', gpv, F_IN, FILL_IN, PCT, RIGHT); repoint(gpname, f"'Waterfall'!$H${r}")
        else:
            S(wf, f'H{r}', gpv, F_LINK, None, PCT, RIGHT)        # Tier 1 %GP = co-invest (link)
        S(wf, f'I{r}', f'=1-H{r}', F_LINK, None, PCT, RIGHT)
        r += 1

    # ---------------------------------------------------------- Waterfall Summary tab
    if 'Waterfall Summary' in wb.sheetnames:
        del wb['Waterfall Summary']
    ws = wb.create_sheet('Waterfall Summary')
    wb.move_sheet('Waterfall Summary', -(len(wb.sheetnames) - 1 - wb.sheetnames.index('Waterfall')))
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '375623'
    ws.column_dimensions['A'].width = 2.5
    ws.column_dimensions['B'].width = 34
    for col in 'CDEFGHIJKLMNOP':
        ws.column_dimensions[col].width = 13

    def tot(row):  # whole-hold sum of a monthly engine row
        return f'SUM(Waterfall!E{row}:NL{row})'

    S(ws, 'B2', 'Equity Waterfall — Summary', F_TITLE)
    S(ws, 'B3', '="LP / GP promote  •  "&PropertyName', Font(name="Calibri", size=10, italic=True, color=GREY))

    # Contribution / Distribution summary
    S(ws, 'B5', 'Contribution / Distribution Summary', F_HDR, FILL_HDR, align=LEFT)
    for col in 'CDE': ws[f'{col}5'].fill = FILL_HDR
    S(ws, 'C6', 'Project', F_HDR, FILL_HDR, align=RIGHT); S(ws, 'D6', 'LP', F_HDR, FILL_HDR, align=RIGHT); S(ws, 'E6', 'GP', F_HDR, FILL_HDR, align=RIGHT)
    S(ws, 'B7', 'Contributions', F_LBL, align=LEFT)
    S(ws, 'C7', f'=-({tot(R["lpcon"])}+{tot(R["gpcon"])})', F_LINK, num=USD, align=RIGHT)
    S(ws, 'D7', f'=-{tot(R["lpcon"])}', F_LINK, num=USD, align=RIGHT)
    S(ws, 'E7', f'=-{tot(R["gpcon"])}', F_LINK, num=USD, align=RIGHT)
    S(ws, 'B8', 'Distributions', F_LBL, align=LEFT)
    S(ws, 'C8', f'={tot(R["lpdist"])}+{tot(R["gpdist"])}', F_LINK, num=USD, align=RIGHT)
    S(ws, 'D8', f'={tot(R["lpdist"])}', F_LINK, num=USD, align=RIGHT)
    S(ws, 'E8', f'={tot(R["gpdist"])}', F_LINK, num=USD, align=RIGHT)
    S(ws, 'B9', 'Net Proceeds', F_TOT, None, None, LEFT, TOPB)
    for col in 'CDE':
        S(ws, f'{col}9', f'={col}8+{col}7', F_TOT, None, USD, RIGHT, TOPB)
    S(ws, 'B10', 'IRR', F_LBL, align=LEFT)
    S(ws, 'C10', '=LevIRR', F_LINK, num=PCT, align=RIGHT)
    S(ws, 'D10', '=LP_IRR', F_LINK, num=PCT, align=RIGHT)
    S(ws, 'E10', '=GP_IRR', F_LINK, num=PCT, align=RIGHT)
    S(ws, 'B11', 'MOIC', F_LBL, align=LEFT)
    S(ws, 'C11', f'=IF(C7=0,0,-C8/C7)', F_LINK, num=MULT, align=RIGHT)
    S(ws, 'D11', '=Waterfall!D16', F_LINK, num=MULT, align=RIGHT)
    S(ws, 'E11', '=Waterfall!D17', F_LINK, num=MULT, align=RIGHT)

    # Tier table
    S(ws, 'B14', 'Distributions by Tier', F_HDR, FILL_HDR, align=LEFT)
    for col in 'CDEFG': ws[f'{col}14'].fill = FILL_HDR
    S(ws, 'C15', 'LP', F_HDR, FILL_HDR, align=RIGHT); S(ws, 'D15', 'LP %', F_HDR, FILL_HDR, align=RIGHT)
    S(ws, 'E15', 'GP', F_HDR, FILL_HDR, align=RIGHT); S(ws, 'F15', 'GP %', F_HDR, FILL_HDR, align=RIGHT)
    S(ws, 'G15', 'Total at Tier', F_HDR, FILL_HDR, align=RIGHT)
    tierdefs = [('Tier 1 — RoC + Pref', R['tier1'], 'GPCoInvest'),
                ('Tier 2 — Pref to Hurdle 2', R['tier2'], 'Promote2'),
                ('Tier 3 — Hurdle 2 to 3', R['tier3'], 'Promote3'),
                ('Tier 4 — Residual', R['tier4'], 'Promote4')]
    rr = 16
    for lab, row, gpshare in tierdefs:
        S(ws, f'B{rr}', lab, F_LBL, align=LEFT)
        S(ws, f'E{rr}', f'={tot(row)}*{gpshare}', F_LINK, num=USD, align=RIGHT)
        S(ws, f'C{rr}', f'={tot(row)}*(1-{gpshare})', F_LINK, num=USD, align=RIGHT)
        S(ws, f'G{rr}', f'={tot(row)}', F_LINK, num=USD, align=RIGHT)
        S(ws, f'D{rr}', f'=IF(G{rr}=0,0,C{rr}/G{rr})', F_LINK, num=PCT, align=RIGHT)
        S(ws, f'F{rr}', f'=IF(G{rr}=0,0,E{rr}/G{rr})', F_LINK, num=PCT, align=RIGHT)
        rr += 1
    S(ws, f'B{rr}', 'Total', F_TOT, None, None, LEFT, TOPB)
    S(ws, f'C{rr}', f'=SUM(C16:C{rr-1})', F_TOT, None, USD, RIGHT, TOPB)
    S(ws, f'E{rr}', f'=SUM(E16:E{rr-1})', F_TOT, None, USD, RIGHT, TOPB)
    S(ws, f'G{rr}', f'=SUM(G16:G{rr-1})', F_TOT, None, USD, RIGHT, TOPB)
    S(ws, f'D{rr}', f'=IF(G{rr}=0,0,C{rr}/G{rr})', F_TOT, None, PCT, RIGHT, TOPB)
    S(ws, f'F{rr}', f'=IF(G{rr}=0,0,E{rr}/G{rr})', F_TOT, None, PCT, RIGHT, TOPB)

    # Annual LP / GP net proceeds
    hdr = rr + 2
    S(ws, f'B{hdr}', 'Annual Net Proceeds', F_HDR, FILL_HDR, align=LEFT)
    S(ws, f'B{hdr+1}', 'Year →', F_LBL, align=LEFT)
    years = 15
    for y in range(1, years + 1):
        col = get_column_letter(3 + y - 1)
        S(ws, f'{col}{hdr}', y, F_HDR, FILL_HDR, num='#,##0', align=RIGHT)
    def annual(label, row, rr2, bold=False):
        S(ws, f'B{rr2}', label, F_TOT if bold else F_LBL, align=LEFT)
        for y in range(1, years + 1):
            col = get_column_letter(3 + y - 1)
            f = (f'=SUM(INDEX(Waterfall!$E${row}:$NL${row},(({y}-1)*12+1)):'
                 f'INDEX(Waterfall!$E${row}:$NL${row},(({y}-1)*12+1)+11))')
            S(ws, f'{col}{rr2}', f, F_TOT if bold else F_LBL, num=USD, align=RIGHT)
    annual('LP Net Proceeds', R['lpnet'], hdr + 2)
    annual('GP Net Proceeds', R['gpnet'], hdr + 3)
    annual('GP Net Proceeds incl. Fees', R['gpfee'], hdr + 4, bold=True)
    ws.freeze_panes = 'C7'

    return {'gp_coinvest_cell': 'Waterfall!D6', 'summary_tab': 'Waterfall Summary'}


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    print('applied:', build(wb)); wb.save(dst); print('saved', dst)


if __name__ == '__main__':
    main()
