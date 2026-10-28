"""
santan_budget.py — expand the Development Budget to the full Santan Valley (Wolff) chart
of accounts, then rebuild the DevSpend cost-curve engine row-for-row to match.

The reference model (LB_AZ Pinal County / Hunt Hwy & Ellsworth, 288 du) carries a detailed
institutional budget: Land; Soft Costs split into Due Diligence / Entitlements / Consultants
/ Other Soft / Financing; and Hard Costs.  This script reproduces every cost line item as an
editable budget line, preserving the current Verde totals by mapping the existing populated
values onto their Santan equivalents (same amount + phase + curve, so monthly spend timing,
capitalized interest and TDC are all unchanged) and defaulting every new line to $0.

Design (keeps the model non-circular and the engine intact in shape):
  * Budget tab, direct-cost block (rows 5..LB): Land / Soft / Hard lines.  DirectCosts sums
    this block; the Group column (C) drives the Land/Soft/Hard subtotals.
  * Budget tab, memo block (below the direct block, Group="Memo"): financing costs, reserves,
    construction-period tax and the sponsor development fee — linked to the model's existing
    capital-stack values and EXCLUDED from DirectCosts, so they appear in the budget for
    completeness without being double-counted in TDC (they already reach TDC via the stack).
  * DevSpend engine: one spread row per direct-cost line (same IF/curve formula and NN..NU
    helper columns as before), relocated summary rows (Total Development Spend, Sponsor Dev
    Fee, Construction-Period Tax), and the cost-curve reconciliation checks.
  * Every external reference to the relocated DevSpend summary rows (34/35/38) and the named
    cells (ConTaxTotal, Chk_CostCurve, Chk_DevFeeCurve, BudgetLand/Soft/Hard, DirectCosts,
    DevFee, BudgetTDCref) is repointed.

Sponsor-specific names from the reference ("… - Wolff") are neutralized to "Sponsor".
Lease-up revenue offsets from the reference budget are not cost line items and are left to the
model's Operating/LeaseUp tabs, so they are not added here.
"""
import sys, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string as ci
from openpyxl.workbook.defined_name import DefinedName

NAVY = "FF1F3864"; INK = "FF000000"; BLUE = "FF0000CC"; GREEN = "FF006100"; GREY = "FF595959"
FILL_SECT = PatternFill("solid", fgColor=NAVY)
FILL_SUB = PatternFill("solid", fgColor="FF2E5496")
FILL_IN = PatternFill("solid", fgColor="FFFDF3D0")
FILL_TOT = PatternFill("solid", fgColor="FFE2EFDA")
F_TITLE = Font(name="Calibri", size=15, bold=True, color=NAVY)
F_SECT = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
F_SUB = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")
F_LBL = Font(name="Calibri", size=10, color=INK)
F_IN = Font(name="Calibri", size=10, color=BLUE)
F_LINK = Font(name="Calibri", size=10, color=GREEN)
F_TOT = Font(name="Calibri", size=10, bold=True, color=INK)
F_FORM = Font(name="Calibri", size=10, color=INK)
F_HDR = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")
F_HINT = Font(name="Calibri", size=9, italic=True, color=GREY)
RIGHT = Alignment(horizontal="right"); LEFT = Alignment(horizontal="left"); CENTER = Alignment(horizontal="center")
TOPB = Border(top=Side(style="thin", color="FF808080"))
USD = '#,##0;(#,##0);"–"'; PCT = '0.0%'; FIRST, LAST = 5, 376

# ---- the chart of accounts ---------------------------------------------------------------
# direct-cost lines: (name, group, category, basis, amount, phase, curve)
SL = 'Straight_Line'; BC = 'Bell_Curve'; SP = 'Single_Point'
DIRECT = [
    # LAND
    ('Land Purchase Costs',              'Land', 'Land', 'abs',  '=LandCost',            'Initial Closing', SP),
    ('Broker Fee',                       'Land', 'Land', 'abs',  0,                      'Entitlement',     SP),
    ('Other Land Cost',                  'Land', 'Land', 'abs',  0,                      'Entitlement',     SP),
    # SOFT — Due Diligence
    ('Legal: Entity Formation',          'Soft', 'Due Diligence', 'abs', 0,              'Entitlement',     SL),
    ('Legal: Closing Costs/Title/Escrow','Soft', 'Due Diligence', 'abs', 0,              'Entitlement',     SL),
    ('Legal: Misc.',                     'Soft', 'Due Diligence', 'abs', 0,              'Construction',    SL),
    ("HOA/CC&R's/Land Use",              'Soft', 'Due Diligence', 'abs', 0,              'Lease-Up',        SL),
    ('Market Study/Appraisal',           'Soft', 'Due Diligence', 'abs', 0,              'Entitlement',     SL),
    ('Other Due Diligence (DD Report, Tax, Misc.)', 'Soft', 'Due Diligence', 'abs', 60000, 'Due Diligence', SL),
    # SOFT — Entitlements
    ('Municipal Fees: Entitlement Phase','Soft', 'Entitlements', 'abs', 0,               'Entitlement',     SL),
    ('Municipal Fees: Plan Review Fees', 'Soft', 'Entitlements', 'abs', 0,               'Entitlement',     SL),
    ('Municipal Fees: Development Fees', 'Soft', 'Entitlements', 'abs', 460000,          'Entitlement',     SL),
    ('Municipal Fees: Permit Fees',      'Soft', 'Entitlements', 'abs', 0,               'Entitlement',     SL),
    ('Municipal Fees: Use Taxes',        'Soft', 'Entitlements', 'abs', 0,               'Entitlement',     SL),
    ('Guarantees/LOC/Bond Premiums',     'Soft', 'Entitlements', 'abs', 0,               'Pre-Construction',SL),
    ('Legal: Zoning/Entitlement',        'Soft', 'Entitlements', 'abs', 0,               'Entitlement',     SL),
    ('Entitlement Consultants',          'Soft', 'Entitlements', 'abs', 168000,          'Entitlement',     SL),
    ('Public Relations/Outreach',        'Soft', 'Entitlements', 'abs', 0,               'Construction',    SL),
    ('Other Entitlement Costs',          'Soft', 'Entitlements', 'abs', 0,               'Entitlement',     SL),
    # SOFT — Consultants
    ('Site Planning',                    'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Architect 1',                      'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Architect 2',                      'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Design Consulting',                'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Structural Engineer',              'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Mechanical Engineer',              'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Plumbing Engineer',                'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Electrical Engineer',              'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Energy Calcs',                     'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Civil Engineer',                   'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Landscape Architect',              'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Dry Utilities',                    'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Fire Sprinkler',                   'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Fire Alarm',                       'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Low Voltage',                      'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Interior Design',                  'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Branding/Graphic Design',          'Soft', 'Consultants', 'abs', 0, 'Lease-Up',        SL),
    ('Signage Design',                   'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Renderings/Photography',           'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Pre-ALTA Survey/Topo',             'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Post ALTA Survey',                 'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Printing/Reimbursable A/E',        'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Traffic Engineer',                 'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Waterproofing/Envelope',           'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Acoustical Engineer',              'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Accessibility Review',             'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Environmental: Phase 1 ESA',       'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Environmental: Phase 2 ESA',       'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Environmental: Geotechnical/Soils','Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Environmental: Asbestos/Lead',     'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Environmental: Wetland/Biology',   'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('Affordable Housing',               'Soft', 'Consultants', 'abs', 0, 'Entitlement',     SL),
    ('S.W.P.P.',                         'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Special Inspections/Testing',      'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Preconstruction Estimating',       'Soft', 'Consultants', 'abs', 0, 'Pre-Construction',SL),
    ('LEED Fees',                        'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Consultant Contingency',           'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Building Maintenance Support',     'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Retail Consulting',                'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Site Security',                    'Soft', 'Consultants', 'abs', 0, 'Construction',    SL),
    ('Surveying: Legal Description, Plat, Mapping', 'Soft', 'Consultants', 'abs', 0, 'Entitlement', SL),
    # SOFT — Other Soft Costs
    ('Property Insurance',               'Soft', 'Other Soft', 'abs', 0, 'Construction', SL),
    ('Working Capital',                  'Soft', 'Other Soft', 'abs', 0, 'Entitlement',  SL),
    ('Sales and Marketing',              'Soft', 'Other Soft', 'abs', 0, 'Lease-Up',     SL),
    ('FF&E: Exercise Equipment',         'Soft', 'Other Soft', 'abs', 0, 'Lease-Up',     SL),
    ('FF&E: Model Furniture',            'Soft', 'Other Soft', 'abs', 0, 'Lease-Up',     SL),
    ('FF&E: Clubhouse Furniture',        'Soft', 'Other Soft', 'abs', 0, 'Lease-Up',     SL),
    ('FF&E: Pool & Exterior Common Areas','Soft', 'Other Soft', 'abs', 0, 'Lease-Up',    SL),
    ('Overhead Reimbursement - Staff',   'Soft', 'Other Soft', 'abs', 0, 'Construction', SL),
    ('Overhead Reimbursement - Remote Office', 'Soft', 'Other Soft', 'abs', 0, 'Construction', SL),
    ('Overhead Reimbursement - Sponsor', 'Soft', 'Other Soft', 'abs', 0, 'Construction', SL),
    ('Travel',                           'Soft', 'Other Soft', 'abs', 0, 'Construction', SL),
    ('Development Fee - Third Party',    'Soft', 'Other Soft', 'abs', 0, 'Construction', SL),
    ('Retail Leasing Commissions',       'Soft', 'Other Soft', 'abs', 0, 'Lease-Up',     SL),
    ('Soft Cost Contingency',            'Soft', 'Other Soft', 'abs', 0, 'Construction', SL),
    ('Other Soft Cost Item 1',           'Soft', 'Other Soft', 'abs', 0, 'Entitlement',  SL),
    ('Other Soft Cost Item 2',           'Soft', 'Other Soft', 'abs', 0, 'Entitlement',  SL),
    ('Other Soft Cost Item 3',           'Soft', 'Other Soft', 'abs', 0, 'Entitlement',  SL),
    ('Other Soft Cost Item 4',           'Soft', 'Other Soft', 'abs', 0, 'Entitlement',  SL),
    # HARD
    ('GC Contract',                      'Hard', 'Hard', 'unit', 206624, 'Construction', BC),
    ('Off Site Improvements (Not Included Above)', 'Hard', 'Hard', 'abs', 0, 'Construction', SL),
    ('Demolition (Not Included Above)',  'Hard', 'Hard', 'abs', '=ExistingStructureCost', 'Construction', SL),
    ('Utility System Installation Costs','Hard', 'Hard', 'abs', 0, 'Construction', SL),
    ('Other Hard Costs Outside GC Contract', 'Hard', 'Hard', 'abs', 0, 'Lease-Up', SL),
    ('Hard Cost Contingency',            'Hard', 'Hard', 'abs', 0, 'Construction', SL),
    ('Tenant Improvements',              'Hard', 'Hard', 'abs', 0, 'Lease-Up',     SL),
    ('Taxes - State Excise & Other',     'Hard', 'Hard', 'abs', 0, 'Construction', BC),
    ('Insurance - Builders Risk, OCIP, Pollution', 'Hard', 'Hard', 'abs', 0, 'Construction', SL),
    ('Unit Finish Upgrades Outside GC Contract', 'Hard', 'Hard', 'abs', 0, 'Construction', SL),
    ('GC Fee',                           'Hard', 'Hard', 'abs', 0, 'Construction', BC),
    ('Other Hard Cost Item 1',           'Hard', 'Hard', 'abs', 0, 'Construction', SL),
    ('Other Hard Cost Item 2',           'Hard', 'Hard', 'abs', 0, 'Construction', SL),
    ('Other Hard Cost Item 3',           'Hard', 'Hard', 'abs', 0, 'Construction', SL),
    ('Other Hard Cost Item 4',           'Hard', 'Hard', 'abs', 0, 'Construction', SL),
]
# memo lines (linked to existing capital-stack values; excluded from DirectCosts)
MEMO = [
    ('Points / Origination',                      '=ConPoints'),
    ('Other Loan / Closing Costs',                '=ConOtherCosts'),
    ('Recourse / Indemnity Fee',                  '=RecourseFee'),
    ('Debt Service (Interest Reserve)',           '=InterestReserveFunded'),
    ('Operating Reserve / Initial Operating Deficit', '=OpReserve'),
    ('Construction-Period Property Taxes',        '=ConTaxTotal'),
    ('Sponsor Development Fee',                    '=DevFee'),
]
CATS = ['Land', 'Due Diligence', 'Entitlements', 'Consultants', 'Other Soft', 'Hard']


def build(wb):
    bg = wb['Budget']; ds = wb['DevSpend']
    Nd = len(DIRECT)
    LB = 4 + Nd                      # last Budget direct row (rows 5..LB)
    LBDS = 5 + Nd                    # last DevSpend spread row (rows 6..LBDS)

    def repoint(name, ref):
        if name in wb.defined_names: del wb.defined_names[name]
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # ---------------------------------------------------------------- Budget tab
    # wipe current content rows (keep title rows 1-2, header row 4)
    for r in range(5, bg.max_row + 1):
        for c in range(1, 14):
            bg.cell(r, c).value = None
            bg.cell(r, c).fill = PatternFill(); bg.cell(r, c).border = Border()

    # clarify the Modifier unit: in this model it shifts the phase start/end by whole MONTHS
    # (DevSpend adds it to a milestone month index), unlike the reference model which used days.
    bg['I4'].value = 'Modifier (months)'
    bg['I3'].value = '± months vs. phase'
    bg['I3'].font = F_HINT; bg['I3'].alignment = RIGHT
    bg['B2'].value = ('Basis: abs=$  unit=$/unit  sf=$/NRSF.  Modifier shifts the phase start/end '
                      'by whole months (+later / −earlier).  Modeled Total drives the cost-curve engine.')

    def bset(r, col, v, font=None, fill=None, num=None, align=None, border=None):
        c = bg.cell(r, ci(col)); c.value = v
        if font: c.font = font
        if fill: c.fill = fill
        if num: c.number_format = num
        if align: c.alignment = align
        if border: c.border = border
        return c

    def line_formulas(r):
        bset(r, 'J', f'=IF(E{r}="abs",F{r},IF(E{r}="unit",F{r}*TotalUnits,IF(E{r}="sf",F{r}*NRSF,F{r})))', F_FORM, None, USD, RIGHT)
        bset(r, 'K', f'=IF(TotalUnits=0,0,J{r}/TotalUnits)', F_FORM, None, USD, RIGHT)
        bset(r, 'L', f'=IF(NRSF=0,0,J{r}/NRSF)', F_FORM, None, USD, RIGHT)
        bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_FORM, None, PCT, RIGHT)

    r = 5
    for name, group, cat, basis, amt, phase, curve in DIRECT:
        bset(r, 'B', name, F_LBL, align=LEFT)
        bset(r, 'C', group, F_FORM, align=LEFT)
        bset(r, 'D', cat, F_HINT, align=LEFT)
        bset(r, 'E', basis, F_IN, FILL_IN, align=CENTER)
        bset(r, 'F', amt, F_IN, FILL_IN, USD, RIGHT)
        bset(r, 'G', phase, F_IN, FILL_IN, align=LEFT)
        bset(r, 'H', curve, F_IN, FILL_IN, align=LEFT)
        bset(r, 'I', 0, F_IN, FILL_IN, '0', RIGHT)
        line_formulas(r)
        r += 1
    # memo block
    r += 1
    mh = r
    bset(r, 'B', 'Financing & Reserves (captured in capitalization — not in Direct Costs)', F_SECT, FILL_SECT, align=LEFT)
    for col in 'CDEFGHIJKLM': bg.cell(r, ci(col)).fill = FILL_SECT
    r += 1
    memo_start = r
    for name, link in MEMO:
        bset(r, 'B', name, F_LBL, align=LEFT)
        bset(r, 'C', 'Memo', F_FORM, align=LEFT)
        bset(r, 'D', 'Financing & Reserves', F_HINT, align=LEFT)
        bset(r, 'E', 'abs', F_FORM, align=CENTER)
        bset(r, 'F', link, F_LINK, None, USD, RIGHT)
        bset(r, 'J', f'=F{r}', F_LINK, None, USD, RIGHT)
        bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_FORM, None, PCT, RIGHT)
        r += 1
    memo_end = r - 1

    # summary block
    r += 1
    bset(r, 'B', 'Budget Summary', F_SECT, FILL_SECT, align=LEFT)
    for col in 'CDEFGHIJKLM': bg.cell(r, ci(col)).fill = FILL_SECT
    r += 1
    land_row = r; bset(r, 'B', 'Land', F_LBL, align=LEFT)
    bset(r, 'J', f'=SUMIF($C$5:$C${LB},"Land",$J$5:$J${LB})', F_FORM, None, USD, RIGHT)
    bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_FORM, None, PCT, RIGHT); r += 1
    soft_row = r; bset(r, 'B', 'Soft Costs', F_LBL, align=LEFT)
    bset(r, 'J', f'=SUMIF($C$5:$C${LB},"Soft",$J$5:$J${LB})', F_FORM, None, USD, RIGHT)
    bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_FORM, None, PCT, RIGHT); r += 1
    hard_row = r; bset(r, 'B', 'Hard Costs', F_LBL, align=LEFT)
    bset(r, 'J', f'=SUMIF($C$5:$C${LB},"Hard",$J$5:$J${LB})', F_FORM, None, USD, RIGHT)
    bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_FORM, None, PCT, RIGHT); r += 1
    direct_row = r; bset(r, 'B', 'Direct Costs (subtotal)', F_TOT, align=LEFT, border=TOPB)
    bset(r, 'J', f'=SUM($J$5:$J${LB})', F_TOT, None, USD, RIGHT, TOPB)
    bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_TOT, None, PCT, RIGHT); r += 1
    devfee_row = r; bset(r, 'B', 'Sponsor Development Fee', F_LBL, align=LEFT)
    bset(r, 'J', '=DevFeePctDirect*DirectCosts', F_FORM, None, USD, RIGHT)
    bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_FORM, None, PCT, RIGHT); r += 1
    tdc_row = r; bset(r, 'B', 'Total Development Cost (TDC)', F_TOT, align=LEFT, border=TOPB)
    bset(r, 'J', '=TDC_Total', F_TOT, None, USD, RIGHT, TOPB)
    bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_TOT, None, PCT, RIGHT); r += 1

    # repoint Budget names
    repoint('BudgetLand', f"Budget!$J${land_row}")
    repoint('BudgetSoft', f"Budget!$J${soft_row}")
    repoint('BudgetHard', f"Budget!$J${hard_row}")
    repoint('DirectCosts', f"Budget!$J${direct_row}")
    repoint('DevFee', f"Budget!$J${devfee_row}")
    repoint('BudgetTDCref', f"Budget!$J${tdc_row}")

    # uses (capitalization) stack
    r += 1
    bset(r, 'B', 'Total Uses (Capitalization)', F_SECT, FILL_SECT, align=LEFT)
    for col in 'CDEFGHIJKLM': bg.cell(r, ci(col)).fill = FILL_SECT
    r += 1
    for lab, f in [('Direct Costs', '=DirectCosts'), ('Development Fee', '=DevFee'),
                   ('Origination Points', '=ConPoints'), ('Recourse / Indemnity Fee', '=RecourseFee'),
                   ('Other Loan / Closing Costs', '=ConOtherCosts'), ('Interest Reserve', '=InterestReserveFunded'),
                   ('Operating Reserve', '=OpReserve'), ('TOTAL USES (TDC)', '=TDC_Total')]:
        tot = lab.startswith('TOTAL')
        bset(r, 'B', lab, F_TOT if tot else F_LBL, align=LEFT, border=TOPB if tot else None)
        bset(r, 'J', f, F_TOT if tot else F_FORM, None, USD, RIGHT, TOPB if tot else None)
        bset(r, 'M', f'=IF(TDC_Total=0,0,J{r}/TDC_Total)', F_TOT if tot else F_FORM, None, PCT, RIGHT)
        r += 1
    r += 1
    bset(r, 'B', 'Sources', F_SECT, FILL_SECT, align=LEFT)
    for col in 'CDEFGHIJKLM': bg.cell(r, ci(col)).fill = FILL_SECT
    r += 1
    for lab, f in [('Construction Loan', '=ConLoanTotal'), ('Equity', '=EquityTotal'),
                   ('TOTAL SOURCES', '=ConLoanTotal+EquityTotal'),
                   ('Sources − Uses (check = 0)', '=ConLoanTotal+EquityTotal-TDC_Total')]:
        tot = lab.startswith('TOTAL') or 'check' in lab
        bset(r, 'B', lab, F_TOT if tot else F_LBL, align=LEFT, border=TOPB if tot else None)
        bset(r, 'J', f, F_TOT if tot else F_FORM, None, USD, RIGHT, TOPB if tot else None)
        r += 1

    # ---------------------------------------------------------------- DevSpend engine
    maxc = ci('NU')
    for rr in range(6, ds.max_row + 1):
        for c in range(2, maxc + 1):
            ds.cell(rr, c).value = None
    NL = get_column_letter(LAST)

    def dset(coord, v, font=None, num=None, align=None):
        c = ds[coord]; c.value = v
        if font: c.font = font
        if num: c.number_format = num
        if align: c.alignment = align

    # spread rows 6..LBDS  (DevSpend row rr <- Budget row rr-1)
    for rr in range(6, LBDS + 1):
        br = rr - 1
        dset(f'B{rr}', f'=Budget!B{br}', F_FORM, align=LEFT)
        dset(f'D{rr}', f'=SUM(E{rr}:{NL}{rr})', F_FORM, USD, RIGHT)
        # helper columns
        dset(f'NM{rr}', f'=Budget!C{br}'); dset(f'NN{rr}', f'=Budget!J{br}')
        dset(f'NO{rr}', f'=Budget!G{br}'); dset(f'NP{rr}', f'=Budget!H{br}')
        dset(f'NQ{rr}', f'=Budget!I{br}')
        dset(f'NR{rr}', f'=MAX(1,MIN(372,INDEX(MilestoneStart,MATCH(NO{rr},MilestoneNames,0))+NQ{rr}))')
        dset(f'NS{rr}', f'=MAX(NR{rr},MIN(372,INDEX(MilestoneEnd,MATCH(NO{rr},MilestoneNames,0))+NQ{rr}))')
        dset(f'NT{rr}', f'=NS{rr}-NR{rr}+1')
        dset(f'NU{rr}', f'=NT{rr}*(NT{rr}+1)*(NT{rr}+2)/6')
        for m in range(1, LAST - FIRST + 2):
            c = get_column_letter(FIRST + m - 1)
            f = (f'=IF(AND({m}>=$NR${rr},{m}<=$NS${rr}),'
                 f'IF($NP${rr}="Single_Point",IF({m}=$NR${rr},$NN${rr},0),'
                 f'IF($NP${rr}="Bell_Curve",$NN${rr}*({m}-$NR${rr}+1)*($NS${rr}-{m}+1)/$NU${rr},'
                 f'$NN${rr}/$NT${rr})),0)')
            dset(f'{c}{rr}', f, F_FORM, USD, RIGHT)

    T = LBDS + 2          # Total Development Spend
    DF = T + 1            # Sponsor Development Fee
    H = T + 2             # Hard Cost Spend
    CH = T + 3            # Cumulative Hard Cost in Place
    TX = T + 4            # Construction-Period Property Tax
    TT = T + 5            # Construction-Period Tax — Total
    CC = T + 7            # cost-curve check
    DFC = T + 8           # dev-fee curve check

    dset(f'B{T}', 'Total Development Spend', F_TOT, align=LEFT); dset(f'D{T}', f'=SUM(E{T}:{NL}{T})', F_TOT, USD, RIGHT)
    dset(f'B{DF}', 'Sponsor Development Fee', F_TOT, align=LEFT); dset(f'D{DF}', f'=SUM(E{DF}:{NL}{DF})', F_TOT, USD, RIGHT)
    dset(f'B{H}', 'Hard Cost Spend', F_FORM, align=LEFT); dset(f'D{H}', f'=SUM(E{H}:{NL}{H})', F_FORM, USD, RIGHT)
    dset(f'B{CH}', 'Cumulative Hard Cost in Place', F_FORM, align=LEFT)
    dset(f'B{TX}', 'Construction-Period Property Tax', F_FORM, align=LEFT); dset(f'D{TX}', f'=SUM(E{TX}:{NL}{TX})', F_FORM, USD, RIGHT)
    dset(f'B{TT}', 'Construction-Period Tax — Total', F_FORM, align=LEFT); dset(f'D{TT}', f'=DevSpend!$D${TX}', F_FORM, USD, RIGHT)
    for m in range(1, LAST - FIRST + 2):
        c = get_column_letter(FIRST + m - 1); p = get_column_letter(FIRST + m - 2)
        dset(f'{c}{T}', f'=SUM({c}6:{c}{LBDS})', F_TOT, USD, RIGHT)
        dset(f'{c}{DF}', (f'=IF(AND({m}>=MS_Initial_Closing_Start,{m}<=MS_Stabilization_Start),'
                          f'DevFee/(MS_Stabilization_Start-MS_Initial_Closing_Start+1),0)'), F_FORM, USD, RIGHT)
        dset(f'{c}{H}', f'=SUMPRODUCT(($NM$6:$NM${LBDS}="Hard")*({c}$6:{c}${LBDS}))', F_FORM, USD, RIGHT)
        dset(f'{c}{CH}', (f'={c}${H}' if m == 1 else f'={p}${CH}+{c}${H}'), F_FORM, USD, RIGHT)
        dset(f'{c}{TX}', (f'=IF(AND({m}>=MS_Initial_Closing_Start,{m}<TaxOpStartMonth),'
                          f'(LandCost+{c}${CH})*EffTaxRateOnValue/12,0)'), F_FORM, USD, RIGHT)

    # reconciliation checks
    dset(f'B{CC-1}', 'Cost-Curve Reconciliation', F_SECT, align=LEFT)
    dset(f'B{CC}', 'Max |line curve sum − Budget total|', F_FORM, align=LEFT)
    terms = ','.join(f'ABS(DevSpend!$D${6+i}-Budget!J{5+i})' for i in range(Nd))
    dset(f'D{CC}', f'=MAX({terms})', F_FORM, USD, RIGHT)
    dset(f'B{DFC}', 'Dev fee curve sum − DevFee', F_FORM, align=LEFT)
    dset(f'D{DFC}', f'=DevSpend!$D${DF}-DevFee', F_FORM, USD, RIGHT)
    repoint('ConTaxTotal', f"DevSpend!$D${TT}")
    repoint('Chk_CostCurve', f"DevSpend!$D${CC}")
    repoint('Chk_DevFeeCurve', f"DevSpend!$D${DFC}")

    # ---------------------------------------------------------------- repoint external refs
    # Match a whole DevSpend reference (single cell OR range), then remap only the row number
    # of each endpoint — so range ends like "$NL$34" in "DevSpend!$E$34:$NL$34" are caught too.
    rowmap = {'34': str(T), '35': str(DF), '38': str(TX)}
    ref_pat = re.compile(r'DevSpend!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?')
    cell_pat = re.compile(r'(\$?[A-Z]{1,3}\$?)(\d+)')
    def fix_ref(mobj):
        return cell_pat.sub(lambda cm: cm.group(1) + rowmap.get(cm.group(2), cm.group(2)), mobj.group(0))
    changed = 0
    for ws in wb.worksheets:
        if ws.title == 'DevSpend':
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'DevSpend!' in c.value:
                    nv = ref_pat.sub(fix_ref, c.value)
                    if nv != c.value:
                        c.value = nv; changed += 1

    # ---------------------------------------------------------------- fix Diagnostics error-scan coverage
    # (a) the DevSpend D-column line-total scan became a reversed range after repoint
    #     (old $D$34:$D$43 -> $D$<T>:$D$43); point it at the relocated line/summary rows.
    # (b) the Budget scan predates the expansion (J5:J39 / D5:D32); widen to the new layout.
    # (c) the Monthly CF and Waterfall tabs were never covered by the in-workbook scan; add them.
    dg = wb['Diagnostics']
    dg['F5'].value = re.sub(r'DevSpend!\$?D\$?\d+:\$?D\$?43', f'DevSpend!$D$6:$D${DFC}', dg['F5'].value)
    dg['F15'].value = ("=SUMPRODUCT(--ISERROR(Budget!J5:J114))"
                       "+SUMPRODUCT(--ISERROR('Monthly CF'!E6:NL41))"
                       "+SUMPRODUCT(--ISERROR(Waterfall!E22:NL48))")
    dg['G15'].value = 'Budget / Monthly CF / Waterfall'

    return {'direct_lines': Nd, 'memo_lines': len(MEMO), 'budget_last_row': LB,
            'devspend_total_row': T, 'devfee_row': DF, 'tax_row': TX, 'ext_cells_repointed': changed,
            'diag_f5': dg['F5'].value[-40:]}


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    print('applied:', build(wb)); wb.save(dst); print('saved', dst)


if __name__ == '__main__':
    main()
