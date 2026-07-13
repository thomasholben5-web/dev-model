"""
Enhancement generator for the Multifamily Ground-Up Development Model.

Operates on an existing, user-edited workbook (single source of truth) and applies:
  Phase A  — Operations fix: lease-up / stabilization milestone reconciles to the
             absorption schedule (resolves the failing diagnostic).
  Phase B  — New "Proforma" tab: institutional Stabilized Year-One Operating
             Proforma with $/SF, $/Unit/Mo, $/Month, $/Unit/Yr, $/Year columns
             and a driver ("% of / amt") column, formatted like a lender proforma.
  Phase C  — Formatting / navigation polish: tab colours, frozen label panes,
             consistent number formats.

Every added value is a formula that references existing named ranges or computed
cells — no hard-keyed calculated outputs are introduced.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "Claude_Dev_Model.xlsx"
DST = sys.argv[2] if len(sys.argv) > 2 else "Claude_Dev_Model_enhanced.xlsx"

# ---- house style (mirrors the existing workbook conventions) -----------------
NAVY = "FF1F3864"
INK = "FF000000"
BLUE_IN = "FF0000CC"      # hard input font
GREEN_LK = "FF006100"     # link font
FILL_IN = "FFFDF3D0"      # input fill (cream)
FILL_TOT = "FFE2EFDA"     # total fill (light green)
FILL_SUB = "FFF2F2F2"     # subtotal / band fill (light grey)

F_TITLE = Font(name="Calibri", size=16, bold=True, color=NAVY)
F_SUB = Font(name="Calibri", size=10, italic=True, color="FF595959")
F_SECT = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
F_LBL = Font(name="Calibri", size=10, color=INK)
F_LBL_B = Font(name="Calibri", size=10, bold=True, color=INK)
F_LINK = Font(name="Calibri", size=10, color=GREEN_LK)
F_FORM = Font(name="Calibri", size=10, color=INK)
F_TOT = Font(name="Calibri", size=10, bold=True, color=INK)
F_HDR = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")

FILL_SECT = PatternFill("solid", fgColor=NAVY)
FILL_TOTAL = PatternFill("solid", fgColor=FILL_TOT)
FILL_BAND = PatternFill("solid", fgColor=FILL_SUB)
FILL_HDR = PatternFill("solid", fgColor="FF2E5496")

RIGHT = Alignment(horizontal="right")
LEFT = Alignment(horizontal="left")
CENTER = Alignment(horizontal="center")
thin = Side(style="thin", color="FFBFBFBF")
TOPB = Border(top=Side(style="thin", color="FF808080"))
UNDER = Border(bottom=thin)

FMT_MONEY = '#,##0;(#,##0)'
FMT_MONEY0 = '#,##0;(#,##0)'
FMT_PSF = '$#,##0.00;($#,##0.00)'
FMT_UMO = '#,##0;(#,##0)'
FMT_PCT = '0.00%'
FMT_AMT = '$#,##0'
FMT_INT = '#,##0'


def style(ws, coord, value=None, font=None, fill=None, num=None, align=None, border=None):
    c = ws[coord]
    if value is not None:
        c.value = value
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if num:
        c.number_format = num
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


# ======================================================================== load
wb = openpyxl.load_workbook(SRC)

# ============================================================ Phase A: fix clock
# Lease-Up milestone duration must equal the number of whole months required for
# cumulative leased units to reach the stabilized target under BOTH the delivery
# and absorption caps. Previously the duration divided the target by the lease
# pace only and produced a fractional month, so INDEX(...,StabMonth) truncated to
# a month at which lease-up was not yet complete -> "reaches stabilized occupancy"
# diagnostic failed. This closed-form matches the LeaseUp engine's own MIN() caps.
tl = wb['Timeline']
tl['D11'].value = ('=MAX(CEILING(StabLeasedTarget/UnitsLeasedPerMo,1),'
                   'CEILING(StabLeasedTarget/UnitsDeliveredPerMo,1))')
tl['D11'].font = F_FORM
tl['D11'].number_format = FMT_INT

# ================================================= Phase B: Stabilized Proforma
if 'Proforma' in wb.sheetnames:
    del wb['Proforma']
ws = wb.create_sheet('Proforma')

for col, w in {'A': 3.7, 'B': 34, 'C': 15, 'D': 12, 'E': 13, 'F': 14, 'G': 13, 'H': 15}.items():
    ws.column_dimensions[col].width = w

style(ws, 'B1', 'Stabilized Year One Operating Proforma', F_TITLE)
style(ws, 'B2',
      '=""&PropertyName&"  •  "&City&", "&State&"  •  Forward 12-month stabilized period"',
      F_SUB)

# ---- context band ------------------------------------------------------------
def band(coord, text):
    style(ws, coord, text, F_SECT, FILL_SECT, align=LEFT)

band('B4', 'Stabilized Period Basis')
ctx = [
    ('B5', 'Stabilization Month', 'C5', '=StabMonth', FMT_INT, 'E5', 'Total Units', 'F5', '=TotalUnits', FMT_INT),
    ('B6', 'Forward Window', 'C6', '="Mo "&StabMonth&" – "&(StabMonth+11)', '@', 'E6', 'Net Rentable SF', 'F6', '=NRSF', FMT_INT),
    ('B7', 'Physical Occupancy', 'C7', '=StabilizedOccupancy', FMT_PCT, 'E7', 'Avg Unit SF', 'F7', '=AvgUnitSF', FMT_INT),
    ('B8', 'Window Start', 'C8', '=INDEX(Operating!$E$4:$NL$4,StabMonth)', 'mmm-yyyy', 'E8', 'Entry Cap Rate', 'F8', '=EntryCapRate', FMT_PCT),
]
for b, blab, c, cf, cfmt, e, elab, f, ff, ffmt in ctx:
    style(ws, b, blab, F_LBL, align=LEFT)
    style(ws, c, cf, F_LINK, num=cfmt, align=RIGHT)
    style(ws, e, elab, F_LBL, align=LEFT)
    style(ws, f, ff, F_LINK, num=ffmt, align=RIGHT)

# ---- table header ------------------------------------------------------------
HDR_ROW = 10
headers = ['', '% of / amt', '$ / SF /mo', '$ / Unit /mo', '$ / Month', '$ / Unit /yr', '$ / Year']
for i, h in enumerate(headers):
    col = get_column_letter(2 + i)
    c = style(ws, f'{col}{HDR_ROW}', h, F_HDR, FILL_HDR,
              align=(LEFT if i == 0 else RIGHT))
style(ws, f'B{HDR_ROW}', 'Stabilized Operating Proforma', F_HDR, FILL_HDR, align=LEFT)

# ---- line driver: (label, operating_row OR custom H-formula, driver formula, driver_fmt, kind)
# kind: 'line'=revenue/expense line, 'sub'=subtotal, 'noi'=noi band, 'total', 'cf'
# operating_row given -> H = fwd-12mo sum of that Operating row.
def opsum(r):
    return (f'=SUM(INDEX(Operating!$E${r}:$NL${r},StabMonth):'
            f'INDEX(Operating!$E${r}:$NL${r},StabMonth+11))')

LINES = [
    ('SECT', 'Revenue', None, None, None),
    ('line', 'Gross Potential Rent', opsum(14), '=GrowthRent', FMT_PCT),
    ('line', 'Vacancy Loss', opsum(15), '=-(1-StabilizedOccupancy)', FMT_PCT),
    ('line', 'Gain / (Loss) to Lease', opsum(16), '=LossToLeasePct', FMT_PCT),
    ('line', 'Bad Debt', opsum(17), '=BadDebtPct', FMT_PCT),
    ('line', 'Concessions', opsum(18), '=ConcessionsPct', FMT_PCT),
    ('line', 'Parking & Storage', opsum(19), '=ParkingPerUnitMo', FMT_AMT),
    ('line', 'Utility Reimbursement', opsum(20), '=UtilReimbPerUnitMo', FMT_AMT),
    ('line', 'Other Income', opsum(21), '=OtherIncPerUnitMo', FMT_AMT),
    ('line', 'Retail Income', opsum(22), '=RetailRentMo', FMT_AMT),
    ('sub', 'Effective Gross Income', opsum(23), None, None),
    ('SECT', 'Operating Expenses', None, None, None),
    ('line', 'Property Taxes', opsum(36), '=GrowthTax', FMT_PCT),
    ('line', 'Insurance', opsum(25), '=ExpInsurancePerUnit', FMT_AMT),
    ('line', 'Utilities', opsum(27), '=ExpUtilitiesPerUnit', FMT_AMT),
    ('line', 'Payroll', opsum(31), '=ExpPayrollPerUnit', FMT_AMT),
    ('line', 'Management Fee', opsum(33), '=MgmtFeePct', FMT_PCT),
    ('line', 'Repairs & Maintenance', opsum(28), '=ExpRMPerUnit', FMT_AMT),
    ('line', 'Turnover', opsum(29), '=ExpTurnoverPerUnit', FMT_AMT),
    ('line', 'Contract Services', opsum(30), '=ExpContractPerUnit', FMT_AMT),
    ('line', 'Admin & Legal', opsum(26), '=ExpAdminPerUnit', FMT_AMT),
    ('line', 'Advertising', opsum(32), '=ExpAdvertisingPerUnit', FMT_AMT),
    ('sub', 'Total Operating Expenses', opsum(37), None, None),
    ('noi', 'Net Operating Income', opsum(38), None, None),
    ('SECT', 'Below-the-Line Expenses', None, None, None),
    ('line', 'Capital Reserves', '=CapReservePerUnit*TotalUnits', '=CapReservePerUnit', FMT_AMT),
    ('line', 'Tax Prep & Appraisal', '=TaxPrepPerUnit*TotalUnits', '=TaxPrepPerUnit', FMT_AMT),
    ('line', 'Asset Management Fees', '=AssetMgmtFeePct*$H$__EGI__', '=AssetMgmtFeePct', FMT_PCT),
    ('sub', 'Total Below-the-Line', None, None, None),  # H filled below (sum of 3)
    ('cf', 'Proforma Operating Cash Flow', opsum(42), None, None),
]

r = HDR_ROW + 1
egi_row = None
below_rows = []
below_total_row = None
noi_row = None
cf_row = None
for kind, label, hform, dform, dfmt in LINES:
    if kind == 'SECT':
        band(f'B{r}', label)
        for col in 'CDEFGH':
            ws[f'{col}{r}'].fill = FILL_SECT
        r += 1
        continue
    # label
    lblfont = F_TOT if kind in ('sub', 'noi', 'cf') else F_LBL
    style(ws, f'B{r}', label, lblfont, align=LEFT)
    # driver
    if dform:
        style(ws, f'C{r}', dform, F_LINK, num=dfmt, align=RIGHT)
    # H (annual)
    if label == 'Effective Gross Income':
        egi_row = r
    if label == 'Asset Management Fees':
        hform = hform.replace('__EGI__', str(egi_row))
    if kind == 'noi':
        noi_row = r
    if kind == 'cf':
        cf_row = r
    if label == 'Total Below-the-Line':
        below_total_row = r
        hform = f'=SUM(H{below_rows[0]}:H{below_rows[-1]})'
    if kind == 'line' and label in ('Capital Reserves', 'Tax Prep & Appraisal', 'Asset Management Fees'):
        below_rows.append(r)
    hfont = F_TOT if kind in ('sub', 'noi', 'cf') else F_FORM
    style(ws, f'H{r}', hform, hfont, num=FMT_MONEY, align=RIGHT)
    # derived columns from H
    style(ws, f'D{r}', f'=IF(NRSF=0,0,H{r}/12/NRSF)', hfont, num=FMT_PSF, align=RIGHT)
    style(ws, f'E{r}', f'=IF(TotalUnits=0,0,H{r}/12/TotalUnits)', hfont, num=FMT_UMO, align=RIGHT)
    style(ws, f'F{r}', f'=H{r}/12', hfont, num=FMT_MONEY, align=RIGHT)
    style(ws, f'G{r}', f'=IF(TotalUnits=0,0,H{r}/TotalUnits)', hfont, num=FMT_MONEY, align=RIGHT)
    # banding / borders
    if kind in ('sub', 'noi', 'cf'):
        for col in 'BCDEFGH':
            ws[f'{col}{r}'].fill = FILL_TOTAL if kind != 'noi' else FILL_BAND
            ws[f'{col}{r}'].border = TOPB
    if kind == 'noi':
        for col in 'BCDEFGH':
            ws[f'{col}{r}'].font = F_TOT
    r += 1

last_row = r - 1

# ---- reconciliation block (local PASS/FAIL, also wired into master diagnostics)
r += 1
band(f'B{r}', 'Reconciliation')
for col in 'CDEFGH':
    ws[f'{col}{r}'].fill = FILL_SECT
rec_start = r + 1
recs = [
    ('Operating CF ties to Operating engine',
     f'=H{cf_row}-SUM(INDEX(Operating!$E$42:$NL$42,StabMonth):INDEX(Operating!$E$42:$NL$42,StabMonth+11))'),
    ('Below-the-line split ties to engine',
     f'=H{below_total_row}-SUM(INDEX(Operating!$E$40:$NL$40,StabMonth):INDEX(Operating!$E$40:$NL$40,StabMonth+11))'),
    ('NOI ties to stabilized NOI (fwd 12mo)',
     f'=H{noi_row}-StabNOI'),
]
r = rec_start
for lab, form in recs:
    style(ws, f'B{r}', lab, F_LBL, align=LEFT)
    style(ws, f'C{r}', form, F_FORM, num=FMT_MONEY, align=RIGHT)
    style(ws, f'D{r}', f'=IF(ABS(C{r})<1,"PASS","FAIL")', F_LBL_B, align=CENTER)
    r += 1
rec_end = r - 1

ws.sheet_view.showGridLines = False
ws.freeze_panes = 'C11'

# move Proforma directly after Dashboard
wb.move_sheet('Proforma', -(len(wb.sheetnames) - 2))

# wire proforma reconciliation into the master error count (no row insertion)
dg = wb['Diagnostics']
dg['C28'].value = (f'=COUNTIF(D4:D25,"FAIL")+COUNTIF(Proforma!$D${rec_start}:$D${rec_end},"FAIL")')

# ================================================= Phase C: navigation / colour
TABCOLOR = {
    'Dashboard': '375623', 'Proforma': '375623', 'Returns': '375623',
    'Diagnostics': 'C00000',
    'Inputs': '2E5496', 'UnitMatrix': '2E5496', 'Budget': '2E5496',
    'Timeline': '7F7F7F', 'DevSpend': '7F7F7F', 'ConLoan': '7F7F7F',
    'LeaseUp': '7F7F7F', 'Operating': '7F7F7F', 'PermLoan': '7F7F7F',
    'Disposition': '7F7F7F', 'CashFlow': '7F7F7F',
}
for name, color in TABCOLOR.items():
    if name in wb.sheetnames:
        wb[name].sheet_properties.tabColor = color

# freeze label columns on the wide monthly grids (col E = month 1 everywhere)
for name in ['DevSpend', 'ConLoan', 'LeaseUp', 'Operating', 'PermLoan', 'CashFlow', 'Returns']:
    if name in wb.sheetnames:
        wb[name].freeze_panes = 'E4'

wb.save(DST)
print(f"saved {DST}")
print(f"Proforma: EGI row {egi_row}, NOI row {noi_row}, CF row {cf_row}, "
      f"below-total row {below_total_row}, recs {rec_start}-{rec_end}")
