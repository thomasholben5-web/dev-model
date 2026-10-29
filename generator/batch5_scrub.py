"""
batch5 — formatting scrub of the Inputs tab (the tab that carried the most leftovers from the
earlier restructuring).  A full-workbook detector found no empty section headers on any other
tab; the Timeline stale fills were already fixed in batch 4.  Value-neutral (clears empty
cells, removes stale fills, reformats two KPI display cells).

  * Remove the "Timing Milestone Durations" header (F8) — durations moved to the Timeline.
  * Remove the "Stabilized NOI Assumptions" header + "now entered on the Proforma" note
    (N16/O16/P16/N17) — those inputs moved to the Proforma.
  * Remove the empty navy box (N26/O26/P26).
  * Fix the ###### KPIs: Static Value ($19M in an 8-wide column) -> compact "$X.XM"; TDC/Unit
    (anchored in the 2-wide Q spacer) -> move the display to the wide R column.
"""
import sys
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill, Alignment

NOFILL = PatternFill()
FILL_KPI = PatternFill("solid", fgColor="FFEAF0F8")
RIGHT = Alignment(horizontal="right")


def build(wb):
    ip = wb['Inputs']

    def clear(coord, keep_fill=False):
        c = ip[coord]; c.value = None
        if not keep_fill:
            c.fill = NOFILL

    # remove leftover empty section headers / notes (clear value + navy fill across the bar)
    for coord in ('F8', 'G8', 'H8', 'I8'):          # "Timing Milestone Durations" header bar
        clear(coord)
    for coord in ('N16', 'O16', 'P16'):             # "Stabilized NOI Assumptions" header bar
        clear(coord)
    ip['N17'].value = None                           # "now entered on the Proforma" note
    for coord in ('N26', 'O26', 'P26'):             # empty navy box
        clear(coord)

    # fix ###### KPIs
    ip['H7'].number_format = '$#,##0.0,,"M"'          # Static Value -> "$19.1M" fits the narrow col
    # TDC / Unit: value was anchored in the 2-wide Q spacer -> move the display to wide R
    ip['Q4'].value = None
    r4 = ip['R4']; r4.value = '=TDCperUnit'; r4.number_format = '$#,##0'
    r4.alignment = RIGHT
    r4.fill = FILL_KPI                                 # keep the KPI box's light-blue fill
    r4.font = copy(ip['N4'].font)                      # match the other KPI value fonts

    return {'inputs_scrubbed': True}


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    print('applied:', build(wb)); wb.save(dst); print('saved', dst)


if __name__ == '__main__':
    main()
