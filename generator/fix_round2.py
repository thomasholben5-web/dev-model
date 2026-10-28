"""
fix_round2.py — items 1,3,4 of the round-2 request (item 2 is verified separately):
  #1  move the stabilized-NOI assumptions onto the Proforma tab as blue hard inputs
      (the former green driver links become the single source); remove them from Inputs
  #3  make the untrended property tax (which drives Static Value / YOC / spread) respond
      to the Dispo Tax Driver: Reassessed -> buyer/reassessed rate, Imputed -> in-place
  (#4 exit valuation already forward-12mo + tax-adjusted; verified by recalc, not changed)
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

BLUE = "FF0000CC"; GREEN = "FF006100"; INK = "FF000000"
FILL_IN = PatternFill("solid", fgColor="FFFDF3D0")
F_IN = Font(name="Calibri", size=10, color=BLUE)
F_LBL = Font(name="Calibri", size=10, color=INK)
F_HINT = Font(name="Calibri", size=9, italic=True, color="FF808080")
F_FORM = Font(name="Calibri", size=10, color=INK)
RIGHT = Alignment(horizontal="right"); LEFT = Alignment(horizontal="left")

# name -> (proforma_row, value, numfmt, (lo,hi,whole))
PCT = '0.00%'; USD = '$#,##0;($#,##0)'; INT = '#,##0'
NOI_INPUTS = {
    'LossToLeasePct': (14, 0, PCT, (0, 0.10, False)),
    'BadDebtPct': (15, 0.005, PCT, (0, 0.10, False)),
    'ConcessionsPct': (16, 0, PCT, (0, 0.10, False)),
    'ModelUnits': (17, 1, INT, (0, 1000, True)),
    'ParkingPerUnitMo': (18, 50, USD, (0, 1e5, False)),
    'UtilReimbPerUnitMo': (19, 50, USD, (0, 1e5, False)),
    'OtherIncPerUnitMo': (20, 50, USD, (0, 1e5, False)),
    'ExpInsurancePerUnit': (24, 500, USD, (0, 1e5, False)),
    'ExpUtilitiesPerUnit': (25, 1000, USD, (0, 1e5, False)),
    'ExpPayrollPerUnit': (26, 1544, USD, (0, 1e5, False)),
    'MgmtFeePct': (27, 0.034, PCT, (0, 0.10, False)),
    'ExpRMPerUnit': (28, 365, USD, (0, 1e5, False)),
    'ExpTurnoverPerUnit': (29, 335, USD, (0, 1e5, False)),
    'ExpContractPerUnit': (30, '=441+408', USD, (0, 1e5, False)),
    'ExpAdminPerUnit': (31, 478, USD, (0, 1e5, False)),
    'ExpAdvertisingPerUnit': (32, 200, USD, (0, 1e5, False)),
    'CapReservePerUnit': (37, 300, USD, (0, 1e5, False)),
    'TaxPrepPerUnit': (38, 0, USD, (0, 1e5, False)),
    'AssetMgmtFeePct': (39, 0.01, PCT, (0, 0.10, False)),
}


def build(wb):
    pf = wb['Proforma']; ip = wb['Inputs']; dg = wb['Diagnostics']

    def repoint(name, ref):
        if name in wb.defined_names: del wb.defined_names[name]
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # ---- #1: NOI assumptions become blue inputs on the Proforma -------------
    for nm, (row, val, fmt, (lo, hi, whole)) in NOI_INPUTS.items():
        # clear the old Inputs cell
        old = wb.defined_names[nm]
        for t, c in old.destinations:
            ip[c.replace('$', '')].value = None
            break
        cc = pf[f'C{row}']
        cc.value = val; cc.font = F_IN; cc.fill = FILL_IN
        cc.number_format = fmt; cc.alignment = RIGHT
        dv = DataValidation(type=('whole' if whole else 'decimal'), operator='between',
                            formula1=repr(lo), formula2=repr(hi), allow_blank=False, showErrorMessage=True)
        dv.promptTitle = nm; dv.prompt = f'Range {lo}–{hi}'
        pf.add_data_validation(dv); dv.add(cc)
        repoint(nm, f"'Proforma'!$C${row}")

    # tidy the (now empty) revenue/expense sections on Inputs -> pointer note
    for r in range(16, 40):
        for col in (14, 15, 16):  # N,O,P
            ip.cell(r, col).value = None
    ip['N16'].value = 'Stabilized NOI Assumptions'
    ip['N16'].font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    ip['N16'].fill = PatternFill("solid", fgColor="FF1F3864")
    ip['O16'].fill = PatternFill("solid", fgColor="FF1F3864")
    ip['P16'].fill = PatternFill("solid", fgColor="FF1F3864")
    ip['N17'].value = '→ now entered on the Proforma tab (blue cells)'
    ip['N17'].font = F_HINT; ip['N17'].alignment = LEFT

    # ---- #3: reassess-aware untrended tax rate ------------------------------
    # named helper on the Proforma (reconciliation area)
    pf['B49'].value = 'Untrended Tax Rate (reassess-aware)'; pf['B49'].font = F_LBL; pf['B49'].alignment = LEFT
    pf['C49'].value = '=IF(ReassessOnSale>=1,DispoEffTaxRate,EffTaxRateOnValue)'
    pf['C49'].font = Font(name="Calibri", size=10, color=GREEN); pf['C49'].number_format = '0.000%'; pf['C49'].alignment = RIGHT
    repoint('UntrendedTaxRate', "'Proforma'!$C$49")
    # property-tax line + its driver now use the reassess-aware rate
    pf['G33'].value = ('=IF(TaxYear1Override>0,TaxYear1Override,'
                       '(G22-SUM(G24:G32))/(EntryCapRate+UntrendedTaxRate)*UntrendedTaxRate)')
    pf['C33'].value = '=UntrendedTaxRate'
    # StaticValue (Inputs!S11) already = StabNOI_Untrended/EntryCap -> auto-reflects.

    # ---- keep the Inputs error-scan honest: add the new Proforma inputs -----
    dg['F17'].value = (dg['F17'].value.rstrip()
                       + '+SUMPRODUCT(--ISERROR(Proforma!C12:C40))')
    return {'moved': len(NOI_INPUTS)}


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    print('applied:', build(wb))
    wb.save(dst)
    print('saved', dst)


if __name__ == '__main__':
    main()
