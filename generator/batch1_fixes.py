"""
batch1_fixes.py — five fixes from the review (milestone consolidation handled separately):

  #1 Proforma: make Loss-to-Lease / Bad Debt / Concessions NEGATIVE inputs, consistent with
     Vacancy Loss (all now entered as negatives; the $/year columns multiply straight through).
  #2 Proforma: enter Vacancy directly on the Proforma (negative input); Stabilized Occupancy
     becomes a derived value (= 1 + vacancy) on Inputs so the lease-up engine still reads it.
  #5 Budget: extend the basis / phase / cost-curve dropdowns from row 31 down to row 97 so the
     full expanded chart of accounts has them.
  #6 DevSpend: formatting — hide the internal engine helper columns, style the period header,
     add year column grouping (like Monthly CF), set column widths.
  #7 Growth: give Growth Start Date a real override input (blank = auto = EOMONTH(Inception,1)).

All sign changes are value-neutral: inputs flip sign AND their consumers flip operator, so EGI,
NOI, TDC, IRR are unchanged.  Milestone consolidation (#4) is a separate script.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter, column_index_from_string as ci
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

BLUE = "FF0000CC"; GREEN = "FF006100"; INK = "FF000000"; NAVY = "FF1F3864"
FILL_IN = PatternFill("solid", fgColor="FFFDF3D0")
FILL_HDR = PatternFill("solid", fgColor="FF2E5496")
F_IN = Font(name="Calibri", size=10, color=BLUE)
F_LINK = Font(name="Calibri", size=10, color=GREEN)
F_LBL = Font(name="Calibri", size=10, color=INK)
F_HINT = Font(name="Calibri", size=9, italic=True, color="FF808080")
F_HDR = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")
RIGHT = Alignment(horizontal="right"); LEFT = Alignment(horizontal="left"); CENTER = Alignment(horizontal="center")
PCT = '0.00%'; FIRST, LAST = 5, 376


def build(wb):
    pf = wb['Proforma']; ip = wb['Inputs']; op = wb['Operating']; bg = wb['Budget']; ds = wb['DevSpend']

    def repoint(name, ref):
        if name in wb.defined_names: del wb.defined_names[name]
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # ---- #2 Vacancy as a Proforma input; Stabilized Occupancy derived -----------------------
    cur_occ = ip['C44'].value
    vac = round(-(1 - cur_occ), 6) if isinstance(cur_occ, (int, float)) else -0.06   # preserve 0.94 -> -0.06
    pf['B13'].value = 'Vacancy Loss'
    pf['C13'].value = vac; pf['C13'].font = F_IN; pf['C13'].fill = FILL_IN; pf['C13'].number_format = PCT; pf['C13'].alignment = RIGHT
    pf['G13'].value = '=G12*VacancyPct'
    repoint('VacancyPct', "'Proforma'!$C$13")
    # Stabilized Occupancy now derived from the vacancy input (keeps the lease-up engine intact)
    ip['B44'].value = 'Stabilized Occupancy (= 1 + vacancy, derived)'
    ip['C44'].value = '=1+VacancyPct'; ip['C44'].font = F_LINK

    # ---- #1 Loss-to-Lease / Bad Debt / Concessions become negative inputs -------------------
    # (values preserved in magnitude: bad debt 0.005 -> -0.005; others already 0)
    for row, cur in ((14, pf['C14'].value), (15, pf['C15'].value), (16, pf['C16'].value)):
        v = cur if isinstance(cur, (int, float)) else 0
        pf[f'C{row}'].value = -abs(v)
        pf[f'C{row}'].font = F_IN; pf[f'C{row}'].fill = FILL_IN; pf[f'C{row}'].number_format = PCT; pf[f'C{row}'].alignment = RIGHT
    pf['G14'].value = '=G12*LossToLeasePct'
    pf['G15'].value = '=G12*BadDebtPct'
    pf['G16'].value = '=G12*ConcessionsPct'
    # refresh data validations on C13:C16 to allow negatives
    for dv in list(pf.data_validations.dataValidation):
        if dv.sqref is not None and any(str(dv.sqref).find(f'C{r}') >= 0 for r in (13, 14, 15, 16)):
            pf.data_validations.dataValidation.remove(dv)
    dvneg = DataValidation(type='decimal', operator='between', formula1='-0.5', formula2='0',
                           allow_blank=False, showErrorMessage=True)
    dvneg.prompt = 'Enter as a negative (loss)'; pf.add_data_validation(dvneg)
    for r in (13, 14, 15, 16):
        dvneg.add(pf[f'C{r}'])
    pf['B10'].value = pf['B10'].value  # (keep header)

    # ---- Operating consumers: flip sign convention (neutral) --------------------------------
    for m in range(1, LAST - FIRST + 2):
        c = get_column_letter(FIRST + m - 1)
        op[f'{c}16'].value = f'={c}$14*LossToLeasePct'
        op[f'{c}17'].value = f'={c}$14*BadDebtPct'
        op[f'{c}18'].value = f'={c}$14*ConcessionsPct'
        # row 68 "Static NOI before Tax": the three pcts are subtracted as positives -> add them
        v68 = op[f'{c}68'].value
        if isinstance(v68, str):
            op[f'{c}68'].value = v68.replace('-LossToLeasePct-BadDebtPct-ConcessionsPct',
                                             '+LossToLeasePct+BadDebtPct+ConcessionsPct')

    # ---- #7 Growth Start Date override ------------------------------------------------------
    ip['B36'].value = 'Growth Start Date (effective)'
    ip['C36'].value = '=IF(ISNUMBER(GrowthStartOverride),GrowthStartOverride,EOMONTH(InceptionDate,1))'
    ip['C36'].font = F_LINK
    ip['B37'].value = '  → Override (blank = auto)'; ip['B37'].font = F_HINT; ip['B37'].alignment = LEFT
    ip['C37'].value = None; ip['C37'].font = F_IN; ip['C37'].fill = FILL_IN
    ip['C37'].number_format = 'mmm-yyyy'; ip['C37'].alignment = RIGHT
    repoint('GrowthStartOverride', "'Inputs'!$C$37")

    # ---- #5 extend Budget dropdowns to the full expanded chart (row 97) ---------------------
    keep = []
    for dv in list(bg.data_validations.dataValidation):
        f1 = (dv.formula1 or '')
        if 'abs' in f1 or 'Single_Point' in f1 or 'Initial Closing' in f1:
            bg.data_validations.dataValidation.remove(dv)
    def add_dv(f1, col):
        dv = DataValidation(type='list', formula1=f1, allow_blank=True, showErrorMessage=True)
        bg.add_data_validation(dv); dv.add(f'{col}5:{col}97')
    add_dv('"abs, sf, unit"', 'E')
    add_dv('"Single_Point, Straight_Line, Bell_Curve"', 'H')
    add_dv('"Initial Closing, Pre-Construction, Due Diligence, Entitlement, Lease-Up, Construction"', 'G')

    # ---- #6 DevSpend formatting -------------------------------------------------------------
    ds.sheet_view.showGridLines = False
    ds.column_dimensions['B'].width = 32
    ds.column_dimensions['C'].width = 2
    ds.column_dimensions['D'].width = 15
    for m in range(1, LAST - FIRST + 2):
        ds.column_dimensions[get_column_letter(FIRST + m - 1)].width = 11
    # style the period header row (row 4) + month-start dates (row 5)
    ds['B4'].value = 'Period →'; ds['B4'].font = F_HDR; ds['B4'].fill = FILL_HDR; ds['B4'].alignment = LEFT
    ds['D4'].value = 'Total'; ds['D4'].font = F_HDR; ds['D4'].fill = FILL_HDR; ds['D4'].alignment = RIGHT
    for m in range(1, LAST - FIRST + 2):
        c = get_column_letter(FIRST + m - 1)
        ds[f'{c}4'].value = m; ds[f'{c}4'].font = F_HDR; ds[f'{c}4'].fill = FILL_HDR
        ds[f'{c}4'].number_format = '#,##0'; ds[f'{c}4'].alignment = RIGHT
        ds[f'{c}5'].number_format = 'mmm-yy'
    ds['B5'].value = 'Month Start'; ds['B5'].font = F_LBL
    # hide the internal engine helper columns (NM..NU)
    for col in ('NM', 'NN', 'NO', 'NP', 'NQ', 'NR', 'NS', 'NT', 'NU'):
        ds.column_dimensions[col].hidden = True
    ds['NM4'].value = 'engine helpers →'; ds['NM4'].font = F_HINT
    # year column grouping (collapse to years), like Monthly CF
    ds.sheet_properties.outlinePr.summaryRight = True
    for y in range(1, 32):
        fc = FIRST + (y - 1) * 12
        g0 = get_column_letter(fc + 1); g1 = get_column_letter(min(fc + 11, LAST))
        if fc + 1 <= LAST:
            ds.column_dimensions.group(g0, g1, outline_level=1, hidden=False)
    ds.freeze_panes = 'E6'

    return {'vacancy': vac, 'occ_derived': '=1+VacancyPct'}


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    print('applied:', build(wb)); wb.save(dst); print('saved', dst)


if __name__ == '__main__':
    main()
