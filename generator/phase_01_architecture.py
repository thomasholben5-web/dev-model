"""
Phase 01 — Workbook architecture, input layer, unit matrix, development budget.

Creates the Inputs sheet (every scalar assumption as a blue input with a
workbook named range), the residential/retail Unit Matrix (with derived unit
counts, NRSF and rent roll), and the Development Budget line-item table (each
line carrying basis / phase / curve / modifier and a modeled total), including
the 10+ user-definable "Other Soft/Hard Cost" lines.
"""

from openpyxl.utils import get_column_letter
from common import (write_title, write_section, write_subhead, label, put_input,
                    put_formula, put_text, set_col_widths, make_table,
                    FMT_USD0, FMT_USD0_DOLLAR, FMT_USD2_DOLLAR, FMT_PCT1, FMT_PCT2,
                    FMT_NUM0, FMT_NUM2, FMT_TEXT, FMT_MONTHS, FMT_MULT,
                    font_subhead, font_label, ALIGN_L, ALIGN_R, FILL_TOTAL,
                    FILL_SUBHEAD)
from assumptions import (SECTIONS, UNIT_PLANS, RETAIL_SUITES, all_budget_lines,
                         N_USER_LINES)
from layout import (SH_INPUT, SH_UNIT, SH_BUDGET, set_row, r)


def build(wb, reg):
    _inputs(wb, reg)
    _unit_matrix(wb, reg)
    _budget(wb, reg)


# ---------------------------------------------------------------------------
def _inputs(wb, reg):
    ws = wb[SH_INPUT]
    set_col_widths(ws, {"A": 34, "B": 16, "C": 44, "D": 4, "E": 20})
    ws.sheet_view.showGridLines = False
    row = write_title(ws, 1, "Institutional Multifamily Development Model — Inputs")
    put_text(ws, 2, 1, "Blue = hard-keyed input.  Every value below is a workbook named range.",
             italic=True)
    row = 4
    for title, rows in SECTIONS:
        row = write_section(ws, row, title, span=3)
        for key, value, fmt, note in rows:
            label(ws, row, key_to_label(key), col=1)
            put_input(ws, row, 2, value, fmt=fmt)
            reg.add(key, SH_INPUT, f"B{row}")
            if note:
                n = ws.cell(row=row, column=3, value=note)
                n.font = font_label(italic=True)
                n.alignment = ALIGN_L
            set_row(SH_INPUT, key, row)
            row += 1
        row += 1
    ws.freeze_panes = "A4"


def key_to_label(key):
    """Human label from a camelCase / mixed key."""
    out = []
    for i, ch in enumerate(key):
        if ch.isupper() and i > 0 and not key[i - 1].isupper():
            out.append(" ")
        out.append(ch)
    return "".join(out)


# ---------------------------------------------------------------------------
def _unit_matrix(wb, reg):
    ws = wb[SH_UNIT]
    set_col_widths(ws, {"A": 16, "B": 12, "C": 10, "D": 10, "E": 12, "F": 12,
                        "G": 14, "H": 14, "I": 10, "J": 10})
    ws.sheet_view.showGridLines = False
    write_title(ws, 1, "Unit Matrix — Residential & Retail")

    hdr = ["Plan", "Class", "Units", "Avg SF", "Total SF", "Market Rent",
           "Net Eff. Rent", "Annual Rent", "Rent / SF", "% of Total"]
    hr = 3
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(row=hr, column=j, value=h)
        c.font = font_subhead()
        c.alignment = ALIGN_R if j >= 3 else ALIGN_L
        c.fill = FILL_SUBHEAD
    first = hr + 1
    row = first
    for (plan, cls, units, sf, rent, ne) in UNIT_PLANS:
        put_text(ws, row, 1, plan)
        put_text(ws, row, 2, cls)
        put_input(ws, row, 3, units, fmt=FMT_NUM0)
        put_input(ws, row, 4, sf, fmt=FMT_NUM0)
        put_formula(ws, row, 5, f"C{row}*D{row}", fmt=FMT_NUM0)
        put_input(ws, row, 6, rent, fmt=FMT_USD0_DOLLAR)
        put_input(ws, row, 7, ne, fmt=FMT_USD0_DOLLAR)
        put_formula(ws, row, 8, f"C{row}*F{row}*12", fmt=FMT_USD0)
        put_formula(ws, row, 9, f"IF(D{row}=0,0,F{row}/D{row})", fmt=FMT_USD2_DOLLAR)
        put_formula(ws, row, 10, f"IF(TotalUnits=0,0,C{row}/TotalUnits)", fmt=FMT_PCT1)
        row += 1
    last = row - 1
    # Totals row
    tr = row
    put_text(ws, tr, 1, "Total / Weighted", bold=True)
    for col, ff in ((3, f"SUM(C{first}:C{last})"), (5, f"SUM(E{first}:E{last})"),
                    (8, f"SUM(H{first}:H{last})")):
        c = put_formula(ws, tr, col, ff, fmt=FMT_NUM0, bold=True)
        c.fill = FILL_TOTAL
    row += 2

    # Derived summary block (named ranges)
    write_subhead(ws, row, "Derived Unit Metrics", span=2)
    row += 1
    def named(lblrow_key, lbl, formula, fmt, name):
        label(ws, lblrow_key, lbl, col=1)
        put_formula(ws, lblrow_key, 2, formula, fmt=fmt, link=True)
        reg.add(name, SH_UNIT, f"B{lblrow_key}")
        set_row(SH_UNIT, name, lblrow_key)
    named(row, "Total Units", f"SUM(C{first}:C{last})", FMT_NUM0, "TotalUnits"); row += 1
    named(row, "Market-Rate Units", f'SUMIF(B{first}:B{last},"Market",C{first}:C{last})', FMT_NUM0, "MarketUnits"); row += 1
    named(row, "Low-Income Units", f'SUMIF(B{first}:B{last},"Low",C{first}:C{last})', FMT_NUM0, "LowUnits"); row += 1
    named(row, "Net Rentable SF (NRSF)", f"SUM(E{first}:E{last})", FMT_NUM0, "NRSF"); row += 1
    named(row, "Buildable SF", "NRSF/EfficiencyRatio", FMT_NUM0, "BuildableSF"); row += 1
    named(row, "Avg Unit SF", "IF(TotalUnits=0,0,NRSF/TotalUnits)", FMT_NUM0, "AvgUnitSF"); row += 1
    named(row, "Base Market Rent / mo",
          f'SUMPRODUCT((B{first}:B{last}="Market")*C{first}:C{last}*F{first}:F{last})',
          FMT_USD0, "BaseMarketRentMo"); row += 1
    named(row, "Base Low-Income Rent / mo",
          f'SUMPRODUCT((B{first}:B{last}="Low")*C{first}:C{last}*F{first}:F{last})',
          FMT_USD0, "BaseLowRentMo"); row += 1

    # Retail matrix
    row += 1
    write_subhead(ws, row, "Retail Matrix", span=4); row += 1
    rh = ["Suite", "SF", "Rent / SF / yr", "Annual Rent", "Monthly Rent"]
    for j, h in enumerate(rh, start=1):
        c = ws.cell(row=row, column=j, value=h); c.font = font_subhead()
        c.alignment = ALIGN_R if j >= 2 else ALIGN_L
    row += 1
    rfirst = row
    for (suite, sf, psf) in RETAIL_SUITES:
        put_text(ws, row, 1, suite)
        put_input(ws, row, 2, sf, fmt=FMT_NUM0)
        put_input(ws, row, 3, psf, fmt=FMT_USD2_DOLLAR)
        put_formula(ws, row, 4, f"B{row}*C{row}", fmt=FMT_USD0)
        put_formula(ws, row, 5, f"B{row}*C{row}/12", fmt=FMT_USD0)
        row += 1
    rlast = row - 1
    label(ws, row, "Retail Rent / mo", col=1)
    put_formula(ws, row, 2, f"SUM(E{rfirst}:E{rlast})", fmt=FMT_USD0, link=True)
    reg.add("RetailRentMo", SH_UNIT, f"B{row}")
    set_row(SH_UNIT, "RetailRentMo", row)
    ws.freeze_panes = "A4"


# ---------------------------------------------------------------------------
def _budget(wb, reg):
    ws = wb[SH_BUDGET]
    set_col_widths(ws, {"A": 30, "B": 8, "C": 8, "D": 8, "E": 14, "F": 20,
                        "G": 16, "H": 10, "I": 16, "J": 12, "K": 10})
    ws.sheet_view.showGridLines = False
    write_title(ws, 1, "Development Budget & Cost Assumptions")
    put_text(ws, 2, 1, "Basis: abs=$  unit=$/unit  sf=$/NRSF.  Modeled Total drives the cost-curve engine.",
             italic=True)
    hdr = ["Line Item", "Group", "Type", "Basis", "Amount", "Phase", "Cost Curve",
           "Modifier", "Modeled Total", "Per Unit", "Per SF"]
    hr = 4
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(row=hr, column=j, value=h); c.font = font_subhead()
        c.alignment = ALIGN_R if j >= 5 else ALIGN_L
        c.fill = FILL_SUBHEAD
    lines = all_budget_lines()
    first = hr + 1
    row = first
    for i, ln in enumerate(lines):
        put_text(ws, row, 1, ln["label"])
        put_text(ws, row, 2, ln["group"])
        put_text(ws, row, 3, ln["type"])
        put_input(ws, row, 4, ln["basis"], fmt=FMT_TEXT)
        # Amount: land lines link to named ranges; others are inputs
        if ln["label"] == "Land Acquisition":
            put_formula(ws, row, 5, "LandCost", fmt=FMT_USD0, link=True)
        elif ln["label"] == "Existing Structure/Demo":
            put_formula(ws, row, 5, "ExistingStructureCost", fmt=FMT_USD0, link=True)
        else:
            put_input(ws, row, 5, ln["amount"] or 0, fmt=FMT_USD0 if ln["basis"] == "abs" else FMT_NUM2)
        put_input(ws, row, 6, ln["phase"], fmt=FMT_TEXT)
        put_input(ws, row, 7, ln["curve"], fmt=FMT_TEXT)
        put_input(ws, row, 8, ln["modifier"], fmt=FMT_NUM0)
        # Modeled total by basis
        mt = (f'IF(D{row}="abs",E{row},IF(D{row}="unit",E{row}*TotalUnits,'
              f'IF(D{row}="sf",E{row}*NRSF,E{row})))')
        put_formula(ws, row, 9, mt, fmt=FMT_USD0)
        put_formula(ws, row, 10, f"IF(TotalUnits=0,0,I{row}/TotalUnits)", fmt=FMT_USD0)
        put_formula(ws, row, 11, f"IF(NRSF=0,0,I{row}/NRSF)", fmt=FMT_USD2_DOLLAR)
        set_row(SH_BUDGET, f"line{i}", row)
        set_row(SH_BUDGET, f"line{i}_curve_col", row)  # curve text at col G, mod at H
        row += 1
    last = row - 1
    set_row(SH_BUDGET, "line_first", first)
    set_row(SH_BUDGET, "line_last", last)
    ws._budget_n_lines = len(lines)

    # Totals & group subtotals
    row += 1
    write_subhead(ws, row, "Budget Summary", span=4); row += 1
    def total(lbl, formula, name, fmt=FMT_USD0):
        label(ws, row_ref[0], lbl, col=1, bold=True)
        c = put_formula(ws, row_ref[0], 9, formula, fmt=fmt, bold=True, link=True)
        c.fill = FILL_TOTAL
        reg.add(name, SH_BUDGET, f"I{row_ref[0]}")
        set_row(SH_BUDGET, name, row_ref[0])
        row_ref[0] += 1
    row_ref = [row]
    total("Land", f'SUMIF(B{first}:B{last},"Land",I{first}:I{last})', "BudgetLand")
    total("Soft Costs", f'SUMIF(B{first}:B{last},"Soft",I{first}:I{last})', "BudgetSoft")
    total("Hard Costs", f'SUMIF(B{first}:B{last},"Hard",I{first}:I{last})', "BudgetHard")
    total("Direct Costs (subtotal)", f"SUM(I{first}:I{last})", "DirectCosts")
    total("Sponsor Development Fee", "DevFeePctDirect*DirectCosts", "DevFee")
    total("Total Development Cost (per Diagnostics)", "TDC_Total", "BudgetTDCref")
    ws.freeze_panes = "A5"
