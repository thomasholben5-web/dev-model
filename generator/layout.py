"""
layout.py — shared build-time state and grid helpers for the phase modules.

Holds the row registry (so any phase can reference a row another phase built),
the canonical sheet-name constants, and the monthly-grid writer.  Keeps the
physical layout in one place so cross-sheet formulas never hardcode addresses.
"""

from openpyxl.utils import get_column_letter
from common import (mcol, mcl, N_MONTHS, FIRST_MONTH_COL,
                    font_label, font_formula, font_link, font_subhead,
                    FMT_NUM0, ALIGN_L, ALIGN_R, FILL_TOTAL, FILL_BAND)

# Canonical tab names ------------------------------------------------------
SH_DASH   = "Dashboard"
SH_INPUT  = "Inputs"
SH_UNIT   = "UnitMatrix"
SH_BUDGET = "Budget"
SH_TIME   = "Timeline"
SH_SPEND  = "DevSpend"
SH_CON    = "ConLoan"
SH_LEASE  = "LeaseUp"
SH_OPS    = "Operating"
SH_PERM   = "PermLoan"
SH_DISPO  = "Disposition"
SH_CF     = "CashFlow"
SH_RET    = "Returns"
SH_DIAG   = "Diagnostics"

SHEET_ORDER = [SH_DASH, SH_INPUT, SH_UNIT, SH_BUDGET, SH_TIME, SH_SPEND,
               SH_CON, SH_LEASE, SH_OPS, SH_PERM, SH_DISPO, SH_CF, SH_RET, SH_DIAG]

# Row registry: ROWS[sheet][key] = row number ------------------------------
ROWS = {s: {} for s in SHEET_ORDER}


def set_row(sheet, key, row):
    ROWS[sheet][key] = row
    return row


def r(sheet, key):
    return ROWS[sheet][key]


# Cross-sheet reference helpers --------------------------------------------
def mref(sheet, key, p, absrow=True):
    """Reference to month p on a registered monthly row of another sheet."""
    row = ROWS[sheet][key]
    col = mcl(p)
    return f"'{sheet}'!{col}${row}" if absrow else f"'{sheet}'!{col}{row}"


def tref(sheet, key):
    """Reference to the Total column (C) of a registered row."""
    row = ROWS[sheet][key]
    return f"'{sheet}'!$C${row}"


def cref(sheet, key, col_letter):
    """Reference to a specific column on a registered row."""
    row = ROWS[sheet][key]
    return f"'{sheet}'!{col_letter}${row}"


def scref(sheet, key):
    """Reference to a single-value cell stored in col B (scalars on a sheet)."""
    row = ROWS[sheet][key]
    return f"'{sheet}'!$B${row}"


# Monthly grid writer ------------------------------------------------------
LAST_COL_LETTER = get_column_letter(mcol(N_MONTHS))
FIRST_COL_LETTER = get_column_letter(FIRST_MONTH_COL)


def grid_row(ws, row, label, func, fmt=FMT_NUM0, unit="", total="sum",
             indent=0, bold=False, link=False, band=False, key=None, sheet=None):
    """Write a labelled monthly row.

    label  : col-A text
    func   : callable(p, C, Cprev) -> formula expression string (no leading '=')
             or a python number.  C is the current month column letter, Cprev
             the previous month's (None at p==1).
    total  : 'sum' -> =SUM(first..last); None -> blank; or an explicit formula
             string for col C.
    key/sheet: if given, register this row in ROWS[sheet][key].
    """
    lc = ws.cell(row=row, column=1, value=("    " * indent) + label)
    lc.font = font_label(bold=bold)
    lc.alignment = ALIGN_L
    if unit:
        u = ws.cell(row=row, column=2, value=unit)
        u.font = font_label(italic=True)
        u.alignment = ALIGN_L
    fnt = (font_link(bold=bold) if link else font_formula(bold=bold))
    for p in range(1, N_MONTHS + 1):
        C = mcl(p)
        Cprev = mcl(p - 1) if p > 1 else None
        val = func(p, C, Cprev)
        cell = ws.cell(row=row, column=mcol(p))
        if isinstance(val, (int, float)):
            cell.value = val
        else:
            cell.value = "=" + val
        cell.font = fnt
        cell.number_format = fmt
        cell.alignment = ALIGN_R
        if band:
            cell.fill = FILL_BAND
    # Total column C
    if total == "sum":
        f = f"=SUM({FIRST_COL_LETTER}{row}:{LAST_COL_LETTER}{row})"
        tc = ws.cell(row=row, column=3, value=f)
    elif total is None:
        tc = ws.cell(row=row, column=3)
    else:
        tc = ws.cell(row=row, column=3, value=("=" + total if not str(total).startswith("=") else total))
    tc.font = font_formula(bold=True)
    tc.number_format = fmt
    tc.alignment = ALIGN_R
    if total is not None:
        tc.fill = FILL_TOTAL
    if key is not None and sheet is not None:
        set_row(sheet, key, row)
    return row + 1


def period_header(ws, row, first_label="Period →"):
    lc = ws.cell(row=row, column=1, value=first_label)
    lc.font = font_subhead()
    tc = ws.cell(row=row, column=3, value="Total")
    tc.font = font_subhead()
    tc.alignment = ALIGN_R
    for p in range(1, N_MONTHS + 1):
        c = ws.cell(row=row, column=mcol(p), value=p)
        c.font = font_subhead()
        c.alignment = ALIGN_R
        c.number_format = FMT_NUM0
    return row + 1


def date_header(ws, row, inception_name="InceptionDate"):
    """Row of month-start dates derived from InceptionDate via EDATE (non-volatile)."""
    lc = ws.cell(row=row, column=1, value="Month Start")
    lc.font = font_label(italic=True)
    for p in range(1, N_MONTHS + 1):
        c = ws.cell(row=row, column=mcol(p))
        c.value = f"=EDATE({inception_name},{p-1})"
        c.number_format = "mmm-yy"
        c.font = font_label(italic=True)
        c.alignment = ALIGN_R
    return row + 1
