"""
fix_round3.py — item 5:
  - "Monthly CF" tab: one monthly grid consolidating every cash flow — development
    spend, dev fee, construction tax, debt draws/interest/balance, equity, EGI, opex,
    NOI, asset-mgmt fees, reserves, operating CF, unlevered & levered CF, cumulatives —
    with year column grouping.
  - Dev Budget: add a "% of TDC" column and expand the summary into a full uses stack
    (Direct + Dev Fee + Points + Recourse + Other + Interest Reserve + Op Reserve = TDC)
    plus Sources (Construction Loan + Equity) with a Sources = Uses check.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "FF1F3864"; INK = "FF000000"; GREEN = "FF006100"
FILL_SECT = PatternFill("solid", fgColor=NAVY)
FILL_TOT = PatternFill("solid", fgColor="FFE2EFDA")
FILL_HDR = PatternFill("solid", fgColor="FF2E5496")
F_TITLE = Font(name="Calibri", size=15, bold=True, color=NAVY)
F_SECT = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
F_LBL = Font(name="Calibri", size=10, color=INK)
F_TOT = Font(name="Calibri", size=10, bold=True, color=INK)
F_FORM = Font(name="Calibri", size=10, color=INK)
F_HDR = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")
F_LINK = Font(name="Calibri", size=10, color=GREEN)
RIGHT = Alignment(horizontal="right"); LEFT = Alignment(horizontal="left")
TOPB = Border(top=Side(style="thin", color="FF808080"))
MONEY = '#,##0;(#,##0);"–"'
FIRST, LAST = 5, 376


def build(wb):
    # ================================================= Monthly CF tab
    if 'Monthly CF' in wb.sheetnames:
        del wb['Monthly CF']
    ws = wb.create_sheet('Monthly CF')
    wb.move_sheet('Monthly CF', -(len(wb.sheetnames) - wb.sheetnames.index('CashFlow')))
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '7F7F7F'
    ws.column_dimensions['A'].width = 2.5
    ws.column_dimensions['B'].width = 36
    ws.column_dimensions['C'].width = 2
    ws.column_dimensions['D'].width = 15
    for m in range(1, LAST - FIRST + 2):
        ws.column_dimensions[get_column_letter(FIRST + m - 1)].width = 12

    def S(coord, v, font=None, fill=None, num=None, align=None):
        c = ws[coord]; c.value = v
        if font: c.font = font
        if fill: c.fill = fill
        if num: c.number_format = num
        if align: c.alignment = align
        return c

    S('B1', 'Monthly Cash Flow — Consolidated', F_TITLE)
    S('B2', '="All flows by month  •  "&PropertyName&"  •  hold "&HoldPeriodMonths&" months"',
      Font(name="Calibri", size=10, italic=True, color="FF595959"))
    # period header + dates
    S('B3', 'Period →', F_HDR, FILL_HDR, align=LEFT); S('D3', 'Total', F_HDR, FILL_HDR, align=RIGHT)
    S('B4', 'Month Ending', F_LBL, align=LEFT)
    for m in range(1, LAST - FIRST + 2):
        col = get_column_letter(FIRST + m - 1)
        S(f'{col}3', m, F_HDR, FILL_HDR, num='#,##0', align=RIGHT)
        S(f'{col}4', f'=Operating!{col}$4', F_FORM, num='mmm-yy', align=RIGHT)

    # (label, kind, per-month formula template with {c}=col ; or None for section)
    ROWS = [
        ('SECT', 'Development & Uses', None),
        ('row', 'Development Spend', '=-DevSpend!{c}$34'),
        ('row', 'Sponsor Development Fee', '=-DevSpend!{c}$35'),
        ('row', 'Construction-Period Property Tax', '=-DevSpend!{c}$38'),
        ('tot', 'Total Project Uses', '=-(DevSpend!{c}$34+DevSpend!{c}$35+DevSpend!{c}$38)'),
        ('SECT', 'Debt', None),
        ('row', 'Construction Debt Draw', '=ConLoan!{c}$29'),
        ('row', 'Construction Interest — Capitalized', '=-ConLoan!{c}$31'),
        ('row', 'Construction Interest — Current Pay', '=-ConLoan!{c}$50'),
        ('row', 'Permanent Interest', '=-PermLoan!{c}$27'),
        ('row', 'Permanent Principal', '=-PermLoan!{c}$28'),
        ('tot', 'Total Cash Debt Service', '=-(ConLoan!{c}$50+PermLoan!{c}$29)'),
        ('bal', 'Construction Loan Balance (EOM)', '=ConLoan!{c}$32'),
        ('bal', 'Permanent Loan Balance (EOM)', '=PermLoan!{c}$30'),
        ('bal', 'Interest Reserve Balance (EOM)', '=ConLoan!{c}$51'),
        ('SECT', 'Equity', None),
        ('row', 'Equity Contribution', '=ConLoan!{c}$27'),
        ('row', 'Refi Cash Distribution', '=IF({c}$3=RefiMonth,RefiCashOut,0)'),
        ('SECT', 'Operations', None),
        ('row', 'Effective Gross Income', '=Operating!{c}$23'),
        ('row', 'Total Operating Expenses', '=-Operating!{c}$37'),
        ('tot', 'Net Operating Income', '=Operating!{c}$38'),
        ('row', 'Asset Management Fees', '=-AssetMgmtFeePct*Operating!{c}$23'),
        ('row', 'Capital Reserves & Tax Prep', '=-(Operating!{c}$40-AssetMgmtFeePct*Operating!{c}$23)'),
        ('tot', 'Property Operating Cash Flow', '=Operating!{c}$42'),
        ('SECT', 'Project Cash Flow', None),
        ('row', 'Net Sale Proceeds (at disposition)', '=IF({c}$3=DispoMonth,NetSale,0)'),
        ('tot', 'Unlevered Cash Flow', '=CashFlow!{c}$7'),
        ('tot', 'Levered Cash Flow to Equity', '=CashFlow!{c}$6'),
        ('cum', 'Cumulative Levered Cash Flow', None),
        ('cum', 'Cumulative Equity Invested', None),
    ]
    r = 5
    lev_row = None
    for kind, label, tmpl in ROWS:
        if kind == 'SECT':
            S(f'B{r}', label, F_SECT, FILL_SECT, align=LEFT)
            S(f'D{r}', '', F_SECT, FILL_SECT)
            for m in range(1, LAST - FIRST + 2):
                ws[f'{get_column_letter(FIRST+m-1)}{r}'].fill = FILL_SECT
            r += 1
            continue
        bold = kind in ('tot',)
        S(f'B{r}', label, F_TOT if bold else F_LBL, align=LEFT)
        if kind == 'cum':
            # cumulative rows reference the levered CF row
            for m in range(1, LAST - FIRST + 2):
                col = get_column_letter(FIRST + m - 1); prev = get_column_letter(FIRST + m - 2)
                if 'Invested' in label:
                    f = (f'=MIN(0,{"" if m==1 else prev+str(r)})+MIN(0,{get_column_letter(FIRST+m-1)}{lev_row})'
                         if m > 1 else f'=MIN(0,{col}{lev_row})')
                    f = f'={("0" if m==1 else prev+str(r))}+MIN(0,{col}{lev_row})'
                else:
                    f = f'={("0" if m==1 else prev+str(r))}+{col}{lev_row}'
                S(f'{col}{r}', f, F_FORM, num=MONEY, align=RIGHT)
            S(f'D{r}', f'=INDEX(E{r}:NL{r},HoldPeriodMonths)', F_TOT, num=MONEY, align=RIGHT)
        else:
            if label == 'Levered Cash Flow to Equity':
                lev_row = r
            for m in range(1, LAST - FIRST + 2):
                col = get_column_letter(FIRST + m - 1)
                S(f'{col}{r}', tmpl.format(c=col), F_TOT if bold else F_FORM, num=MONEY, align=RIGHT)
            if kind in ('bal',):
                S(f'D{r}', f'=INDEX(E{r}:NL{r},HoldPeriodMonths)', F_FORM, num=MONEY, align=RIGHT)
            else:
                S(f'D{r}', f'=SUM(E{r}:NL{r})', F_TOT if bold else F_FORM, num=MONEY, align=RIGHT)
        if bold:
            for cc in ('B', 'D'): ws[f'{cc}{r}'].border = TOPB
        r += 1

    # year column grouping (collapse to years)
    ws.sheet_properties.outlinePr.summaryRight = True
    for y in range(1, 32):
        fc = FIRST + (y - 1) * 12
        g0 = get_column_letter(fc + 1); g1 = get_column_letter(min(fc + 11, LAST))
        if fc + 1 <= LAST:
            ws.column_dimensions.group(g0, g1, outline_level=1, hidden=False)
    ws.freeze_panes = 'E5'

    # ================================================= Dev Budget enhancements
    bg = wb['Budget']
    # % of TDC column on each line + summary
    bg['M4'].value = '% of TDC'; bg['M4'].font = F_HDR; bg['M4'].fill = FILL_HDR; bg['M4'].alignment = RIGHT
    for r in range(5, 40):
        if bg.cell(r, 10).value is not None:  # J has a value
            c = bg.cell(r, 13, value=f'=IF(TDC_Total=0,0,J{r}/TDC_Total)')
            c.font = F_FORM; c.number_format = '0.0%'; c.alignment = RIGHT
    # expand summary into full uses stack + sources (rows 41+)
    start = 41
    S2 = lambda coord, v, font=None, fill=None, num=None, align=None: _put(bg, coord, v, font, fill, num, align)
    bg[f'B{start}'].value = 'Total Uses (Capitalization)'; bg[f'B{start}'].font = F_SECT; bg[f'B{start}'].fill = FILL_SECT
    for col in ('J', 'K', 'L', 'M'):
        bg[f'{col}{start}'].fill = FILL_SECT
    uses = [
        ('Direct Costs', '=DirectCosts'), ('Development Fee', '=DevFee'),
        ('Origination Points', '=ConPoints'), ('Recourse / Indemnity Fee', '=RecourseFee'),
        ('Other Loan / Closing Costs', '=ConOtherCosts'), ('Interest Reserve', '=InterestReserveFunded'),
        ('Operating Reserve', '=OpReserve'), ('TOTAL USES (TDC)', '=TDC_Total'),
    ]
    rr = start + 1
    for lab, f in uses:
        tot = lab.startswith('TOTAL')
        _put(bg, f'B{rr}', lab, F_TOT if tot else F_LBL, None, None, LEFT)
        _put(bg, f'J{rr}', f, F_TOT if tot else F_FORM, None, MONEY, RIGHT)
        _put(bg, f'M{rr}', f'=IF(TDC_Total=0,0,J{rr}/TDC_Total)', F_TOT if tot else F_FORM, None, '0.0%', RIGHT)
        if tot:
            bg[f'B{rr}'].border = TOPB; bg[f'J{rr}'].border = TOPB
        rr += 1
    rr += 1
    _put(bg, f'B{rr}', 'Sources', F_SECT, FILL_SECT, None, LEFT)
    for col in ('J', 'K', 'L', 'M'): bg[f'{col}{rr}'].fill = FILL_SECT
    rr += 1
    for lab, f in [('Construction Loan', '=ConLoanTotal'), ('Equity', '=EquityTotal'),
                   ('TOTAL SOURCES', '=ConLoanTotal+EquityTotal'),
                   ('Sources − Uses (check = 0)', '=ConLoanTotal+EquityTotal-TDC_Total')]:
        tot = lab.startswith('TOTAL') or 'check' in lab
        _put(bg, f'B{rr}', lab, F_TOT if tot else F_LBL, None, None, LEFT)
        _put(bg, f'J{rr}', f, F_TOT if tot else F_FORM, None, MONEY, RIGHT)
        if tot: bg[f'B{rr}'].border = TOPB; bg[f'J{rr}'].border = TOPB
        rr += 1
    return {'monthly_rows': r - 5}


def _put(ws, coord, v, font=None, fill=None, num=None, align=None):
    c = ws[coord]; c.value = v
    if font: c.font = font
    if fill: c.fill = fill
    if num: c.number_format = num
    if align: c.alignment = align
    return c


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    print('applied:', build(wb)); wb.save(dst); print('saved', dst)


if __name__ == '__main__':
    main()
