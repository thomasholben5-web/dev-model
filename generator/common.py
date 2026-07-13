"""
common.py — Shared infrastructure for the Multifamily Development Model generator.

Central definitions used by every phase module:
  * Monthly horizon and column mapping
  * Colour / font / fill / number-format conventions (blue input, black formula, green link)
  * A workbook-level named-range registry
  * Cell-writing helpers that enforce the styling conventions
  * A helper to build native Excel Tables

No calculation logic lives here — only layout and styling primitives.
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName

# ---------------------------------------------------------------------------
# Horizon
# ---------------------------------------------------------------------------
# The model is fully monthly.  Hold period toggles 1..360.  We additionally
# carry 12 forward months so that a 12-month forward NOI can always be
# computed even at a 360-month hold.
N_HOLD_MAX = 360          # maximum hold in months (disposition toggle upper bound)
N_FWD = 12                # forward months required for exit-value NOI
N_MONTHS = N_HOLD_MAX + N_FWD   # = 372 monthly columns in every time-series grid

# Column mapping for monthly grids.
#   Col A (1): line-item label
#   Col B (2): unit / basis note
#   Col C (3): Total (period sum, where meaningful)
#   Month p (1..N_MONTHS) -> column (FIRST_MONTH_COL + p - 1)
FIRST_MONTH_COL = 4       # month 1 lives in column D

def mcol(p):
    """1-based month index -> column index (integer)."""
    return FIRST_MONTH_COL + p - 1

def mcl(p):
    """1-based month index -> column letter (e.g. 'D')."""
    return get_column_letter(mcol(p))

# ---------------------------------------------------------------------------
# Colour conventions
# ---------------------------------------------------------------------------
BLUE_INPUT = "FF0000CC"     # hard-keyed user inputs
BLACK_FORMULA = "FF000000"  # in-sheet calculations
GREEN_LINK = "FF006100"     # links to other sheets / named ranges
WHITE = "FFFFFFFF"

FILL_HEADER = PatternFill("solid", fgColor="FF1F3864")     # dark navy section headers
FILL_SUBHEAD = PatternFill("solid", fgColor="FFD6DCE4")    # light band sub-headers
FILL_INPUT = PatternFill("solid", fgColor="FFFDF3D0")      # pale yellow input cells
FILL_TOTAL = PatternFill("solid", fgColor="FFE2EFDA")      # pale green totals
FILL_CHECK_OK = PatternFill("solid", fgColor="FFC6EFCE")
FILL_CHECK_BAD = PatternFill("solid", fgColor="FFFFC7CE")
FILL_BAND = PatternFill("solid", fgColor="FFF2F2F2")

THIN = Side(style="thin", color="FFBFBFBF")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BOTTOM = Border(bottom=Side(style="thin", color="FF808080"))
BORDER_TOP = Border(top=Side(style="thin", color="FF808080"))

# ---------------------------------------------------------------------------
# Number formats
# ---------------------------------------------------------------------------
FMT_USD0 = '#,##0;(#,##0)'
FMT_USD2 = '#,##0.00;(#,##0.00)'
FMT_USD0_DOLLAR = '$#,##0;($#,##0)'
FMT_USD2_DOLLAR = '$#,##0.00;($#,##0.00)'
FMT_PCT1 = '0.0%'
FMT_PCT2 = '0.00%'
FMT_PCT0 = '0%'
FMT_NUM0 = '#,##0'
FMT_NUM1 = '#,##0.0'
FMT_NUM2 = '#,##0.00'
FMT_MULT = '0.00"x"'
FMT_BPS = '#,##0" bps"'
FMT_DATE = 'mmm-yyyy'
FMT_MONTHS = '#,##0" mo"'
FMT_TEXT = '@'

# ---------------------------------------------------------------------------
# Fonts
# ---------------------------------------------------------------------------
def font_input(bold=False):
    return Font(name="Calibri", size=10, color=BLUE_INPUT, bold=bold)

def font_formula(bold=False):
    return Font(name="Calibri", size=10, color=BLACK_FORMULA, bold=bold)

def font_link(bold=False):
    return Font(name="Calibri", size=10, color=GREEN_LINK, bold=bold)

def font_label(bold=False, italic=False):
    return Font(name="Calibri", size=10, color=BLACK_FORMULA, bold=bold, italic=italic)

def font_header():
    return Font(name="Calibri", size=11, color=WHITE, bold=True)

def font_subhead():
    return Font(name="Calibri", size=10, color="FF1F3864", bold=True)

def font_title():
    return Font(name="Calibri", size=16, color="FF1F3864", bold=True)

ALIGN_L = Alignment(horizontal="left", vertical="center")
ALIGN_R = Alignment(horizontal="right", vertical="center")
ALIGN_C = Alignment(horizontal="center", vertical="center")
ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


# ---------------------------------------------------------------------------
# Named-range registry
# ---------------------------------------------------------------------------
class NameRegistry:
    """Collects workbook-level named ranges and applies them to the workbook.

    Also serves as a single lookup so phases can reference an input by name
    without knowing its physical cell address.
    """
    def __init__(self):
        self._names = {}     # name -> "Sheet!$A$1"

    def add(self, name, sheet, cell):
        ref = f"'{sheet}'!{_abs(cell)}"
        if name in self._names and self._names[name] != ref:
            raise ValueError(f"Named range '{name}' redefined: {self._names[name]} -> {ref}")
        self._names[name] = ref
        return name

    def ref(self, name):
        if name not in self._names:
            raise KeyError(f"Named range '{name}' not defined")
        return self._names[name]

    def has(self, name):
        return name in self._names

    def apply(self, wb):
        for name, ref in self._names.items():
            wb.defined_names.add(DefinedName(name, attr_text=ref))

    def dump(self):
        return dict(sorted(self._names.items()))


def _abs(cell):
    """Turn 'D5' into '$D$5'; leave ranges like '$D$5:$F$5' alone."""
    if "!" in cell or "$" in cell:
        return cell
    # split letters/digits
    col = "".join(c for c in cell if c.isalpha())
    row = "".join(c for c in cell if c.isdigit())
    return f"${col}${row}"


# ---------------------------------------------------------------------------
# Cell writing helpers
# ---------------------------------------------------------------------------
def set_col_widths(ws, widths):
    """widths: dict {col_letter: width}."""
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def write_title(ws, row, text, span=8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = font_title()
    return row + 1

def write_section(ws, row, text, span=3, col=1):
    """Dark navy section header band."""
    for j in range(col, col + span):
        cell = ws.cell(row=row, column=j)
        cell.fill = FILL_HEADER
    c = ws.cell(row=row, column=col, value=text)
    c.font = font_header()
    c.alignment = ALIGN_L
    return row + 1

def write_subhead(ws, row, text, span=3, col=1):
    for j in range(col, col + span):
        cell = ws.cell(row=row, column=j)
        cell.fill = FILL_SUBHEAD
    c = ws.cell(row=row, column=col, value=text)
    c.font = font_subhead()
    c.alignment = ALIGN_L
    return row + 1

def label(ws, row, text, col=1, bold=False, italic=False, indent=0):
    c = ws.cell(row=row, column=col, value=("    " * indent) + text)
    c.font = font_label(bold=bold, italic=italic)
    c.alignment = ALIGN_L
    return c

def put_input(ws, row, col, value, fmt=FMT_NUM0, note=None, note_col=None):
    """Blue hard-keyed input cell with pale-yellow fill."""
    c = ws.cell(row=row, column=col, value=value)
    c.font = font_input()
    c.number_format = fmt
    c.fill = FILL_INPUT
    c.alignment = ALIGN_R if fmt not in (FMT_TEXT,) else ALIGN_L
    c.border = BORDER_ALL
    if note is not None and note_col is not None:
        n = ws.cell(row=row, column=note_col, value=note)
        n.font = font_label(italic=True)
        n.alignment = ALIGN_L
    return c

def put_formula(ws, row, col, formula, fmt=FMT_NUM0, link=False, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=formula)
    c.font = font_link(bold=bold) if link else font_formula(bold=bold)
    c.number_format = fmt
    c.alignment = ALIGN_R
    if fill is not None:
        c.fill = fill
    return c

def put_text(ws, row, col, text, bold=False, italic=False, align=ALIGN_L, link=False):
    c = ws.cell(row=row, column=col, value=text)
    c.font = font_link(bold=bold) if link else font_label(bold=bold, italic=italic)
    c.alignment = align
    return c


def make_table(ws, name, first_row, first_col, last_row, last_col, style="TableStyleLight9"):
    ref = f"{get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row}"
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name=style, showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)
    return tab


def month_header_row(ws, row, first_label="Period", freeze_note=True):
    """Write the standard monthly header: period numbers 1..N_MONTHS across."""
    lc = ws.cell(row=row, column=1, value=first_label)
    lc.font = font_subhead()
    for p in range(1, N_MONTHS + 1):
        c = ws.cell(row=row, column=mcol(p), value=p)
        c.font = font_subhead()
        c.alignment = ALIGN_R
        c.number_format = FMT_NUM0
    return row


def freeze_grid(ws, row):
    """Freeze panes so labels (cols A-C) and header rows stay visible."""
    ws.freeze_panes = ws.cell(row=row, column=FIRST_MONTH_COL)
