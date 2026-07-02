"""
Phase 03 — Development budget & cost-curve engine.

Spreads every budget line (including the user-definable lines) across its
milestone-relative phase window using its cost curve (Straight_Line, Bell_Curve,
Single_Point, Development_Fee).  Each line's monthly curve sums exactly to its
modeled total (checked in col C vs. the Budget total).  Also spreads the sponsor
development fee straight-line across Initial Closing..Stabilization.
"""

from openpyxl.utils import get_column_letter
from common import (write_title, write_section, label, put_formula, put_text,
                    set_col_widths, FMT_USD0, FMT_NUM0, FMT_TEXT, FMT_MONTHS,
                    font_subhead, font_formula, ALIGN_R, FILL_TOTAL, FILL_CHECK_OK,
                    N_MONTHS, mcol, mcl)
from assumptions import all_budget_lines
from layout import (SH_SPEND, SH_BUDGET, SH_TIME, set_row, r, period_header,
                    grid_row, FIRST_COL_LETTER, LAST_COL_LETTER)

# Helper columns placed to the right of the monthly grid
H0 = mcol(N_MONTHS) + 2
HC = {k: get_column_letter(H0 + i) for i, k in enumerate(
    ["total", "phase", "curve", "mod", "start", "end", "n", "denom"])}


def build(wb, reg):
    ws = wb[SH_SPEND]
    set_col_widths(ws, {"A": 30})
    ws.sheet_view.showGridLines = False
    write_title(ws, 1, "Development Cost-Curve Engine — Monthly Spend by Line")

    # helper header labels
    for k, col in HC.items():
        c = ws.cell(row=3, column=H0 + list(HC).index(k), value=k)
        c.font = font_subhead()

    row = 4
    row = period_header(ws, row)
    lines = all_budget_lines()

    line_rows = []
    for i, ln in enumerate(lines):
        R = row
        brow = r(SH_BUDGET, f"line{i}")
        # helper cells
        _put(ws, R, HC["total"], f"Budget!I{brow}", FMT_USD0)
        _put(ws, R, HC["phase"], f"Budget!F{brow}", FMT_TEXT, link=True)
        _put(ws, R, HC["curve"], f"Budget!G{brow}", FMT_TEXT, link=True)
        _put(ws, R, HC["mod"], f"Budget!H{brow}", FMT_NUM0)
        _put(ws, R, HC["start"],
             f'MAX(1,MIN({N_MONTHS},INDEX(MilestoneStart,MATCH({HC["phase"]}{R},MilestoneNames,0))+{HC["mod"]}{R}))',
             FMT_MONTHS)
        _put(ws, R, HC["end"],
             f'MAX({HC["start"]}{R},MIN({N_MONTHS},INDEX(MilestoneEnd,MATCH({HC["phase"]}{R},MilestoneNames,0))+{HC["mod"]}{R}))',
             FMT_MONTHS)
        _put(ws, R, HC["n"], f'{HC["end"]}{R}-{HC["start"]}{R}+1', FMT_NUM0)
        _put(ws, R, HC["denom"], f'{HC["n"]}{R}*({HC["n"]}{R}+1)*({HC["n"]}{R}+2)/6', FMT_NUM0)

        def make(R):
            s = f"${HC['start']}${R}"; e = f"${HC['end']}${R}"
            tot = f"${HC['total']}${R}"; cv = f"${HC['curve']}${R}"
            nn = f"${HC['n']}${R}"; dn = f"${HC['denom']}${R}"
            return (lambda p, C, Cprev:
                    f'IF(AND({p}>={s},{p}<={e}),'
                    f'IF({cv}="Single_Point",IF({p}={s},{tot},0),'
                    f'IF({cv}="Bell_Curve",{tot}*({p}-{s}+1)*({e}-{p}+1)/{dn},'
                    f'{tot}/{nn})),0)')
        row = grid_row(ws, R, ln["label"], make(R), fmt=FMT_USD0, total="sum",
                       key=f"line_spend{i}", sheet=SH_SPEND)
        line_rows.append(R)

    lr_first, lr_last = line_rows[0], line_rows[-1]

    # Total development spend (direct lines)
    row += 1
    def make_sum():
        return lambda p, C, Cprev: f"SUM({C}{lr_first}:{C}{lr_last})"
    row = grid_row(ws, row, "Total Development Spend", make_sum(), fmt=FMT_USD0,
                   total="sum", bold=True, key="spend_total", sheet=SH_SPEND)

    # Development fee spread (straight-line, Initial Closing..Stabilization)
    def make_fee():
        return (lambda p, C, Cprev:
                f"IF(AND({p}>=MS_Initial_Closing_Start,{p}<=MS_Stabilization_Start),"
                f"DevFee/(MS_Stabilization_Start-MS_Initial_Closing_Start+1),0)")
    row = grid_row(ws, row, "Sponsor Development Fee", make_fee(), fmt=FMT_USD0,
                   total="sum", key="fee_spend", sheet=SH_SPEND)

    # Reconciliation checks: each line's monthly sum vs Budget total
    row += 1
    write_section(ws, row, "Cost-Curve Reconciliation", span=3); row += 1
    label(ws, row, "Max |line curve sum − Budget total|", col=1)
    diff_terms = [
        f"ABS('{SH_SPEND}'!$C${R}-Budget!I{r(SH_BUDGET, f'line{i}')})"
        for i, R in enumerate(line_rows)]
    put_formula(ws, row, 3, "MAX(" + ",".join(diff_terms) + ")",
                fmt=FMT_USD0, bold=True)
    reg.add("Chk_CostCurve", SH_SPEND, f"C{row}")
    set_row(SH_SPEND, "Chk_CostCurve", row)
    row += 1
    label(ws, row, "Dev fee curve sum − DevFee", col=1)
    put_formula(ws, row, 3, f"'{SH_SPEND}'!$C${r(SH_SPEND,'fee_spend')}-DevFee", fmt=FMT_USD0)
    reg.add("Chk_DevFeeCurve", SH_SPEND, f"C{row}")
    set_row(SH_SPEND, "Chk_DevFeeCurve", row)
    ws.freeze_panes = "D5"


def _put(ws, row, col_letter, formula, fmt, link=False):
    from openpyxl.utils import column_index_from_string
    c = ws.cell(row=row, column=column_index_from_string(col_letter))
    c.value = "=" + formula
    c.number_format = fmt
    c.font = font_formula()
    c.alignment = ALIGN_R
    return c
