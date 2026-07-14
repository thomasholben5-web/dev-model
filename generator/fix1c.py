"""
fix1c.py — engine remediation round 1C on Claude_Dev_Model.xlsx.

Items:
 1C.1 remove Dashboard "Year-1 NOI (trended)" display row (keep Year1NOI + Returns!C23)
 1C.2 unwind Scenarios: restore 8 hard inputs, drop scenario names, delete sheet
 1C.3 reconcile timeline (delivery <= con end, stab >= con end) + 2 hard FAIL diagnostics
 1C.4 refi at stabilization: RefiMonth = MAX(Stab, ConEnd)+RefiLag capped at maturity;
      wire Recap milestone; FAIL diagnostic
 1C.5 interest current-pay toggle: capitalize pre-trigger (from reserve), pay current
      post-trigger (levered CF outflow); reserve drawdown + exhaustion diagnostic
 1C.6 size LTC including capitalized interest (closed form, non-circular) + diagnostic
 1C.7 sizing tests on untrended NOI
 1C.8 dynamic payoff index (ConLoan C47)
 1C.9 handled in validation (two-branch proofs)
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

BLUE = "FF0000CC"; INK = "FF000000"; GREEN = "FF006100"; NAVY = "FF1F3864"
FILL_IN = PatternFill("solid", fgColor="FFFDF3D0")
FILL_SECT = PatternFill("solid", fgColor=NAVY)
F_IN = Font(name="Calibri", size=10, color=BLUE)
F_LBL = Font(name="Calibri", size=10, color=INK)
F_LBLB = Font(name="Calibri", size=10, bold=True, color=INK)
F_LINK = Font(name="Calibri", size=10, color=GREEN)
F_FORM = Font(name="Calibri", size=10, color=INK)
F_SECT = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
F_HINT = Font(name="Calibri", size=9, italic=True, color="FF595959")
RIGHT = Alignment(horizontal="right"); LEFT = Alignment(horizontal="left"); CENTER = Alignment(horizontal="center")
FMT_INT = '#,##0'; FMT_MONEY = '#,##0;(#,##0)'; FMT_PCT = '0.00%'
FIRST, LAST = 5, 376  # month cols E..NL


def S(ws, coord, v=None, font=None, fill=None, num=None, align=None):
    c = ws[coord]
    if v is not None: c.value = v
    if font: c.font = font
    if fill: c.fill = fill
    if num: c.number_format = num
    if align: c.alignment = align
    return c


def repoint(wb, name, ref):
    if name in wb.defined_names:
        del wb.defined_names[name]
    wb.defined_names.add(DefinedName(name, attr_text=ref))


def drop(wb, name):
    if name in wb.defined_names:
        del wb.defined_names[name]


def cols():
    for m in range(1, LAST - FIRST + 2):
        yield m, get_column_letter(FIRST + m - 1)


def build(wb):
    ip = wb['Inputs']; tl = wb['Timeline']; cl = wb['ConLoan']; cf = wb['CashFlow']
    dg = wb['Diagnostics']; db = wb['Dashboard']; an = wb['Annual']

    # ============================================================= 1C.2 unwind
    restore = {16: 0.0525, 21: 0.065, 22: 0.03, 24: 0.03, 33: 18, 36: 0.94, 63: 0.057, 73: 0.0625}
    for r, val in restore.items():
        S(ip, f'C{r}', val, F_IN, FILL_IN, num=(FMT_INT if r == 33 else (FMT_PCT if val < 1 else FMT_MONEY)), align=RIGHT)
    drop(wb, 'ScenarioSelect'); drop(wb, 'ScenarioNum')
    db['B2'].value = '=City&", "&State&"   •   "&TotalUnits&" units   •   "&TEXT(NRSF,"#,##0")&" NRSF"'
    an['B2'].value = '="Hold: "&HoldPeriodMonths&" months   •   Disposition month "&DispoMonth'
    if 'Scenarios' in wb.sheetnames:
        del wb['Scenarios']

    # ===================================================== 1C.1 Dashboard row
    # remove left-column "Year-1 NOI (trended)" (row 15); reflow rows 17..28 up by 1
    for cc in ('B', 'C'):
        db[f'{cc}15'].value = None
    import copy
    for r in range(17, 29):
        for cc in ('B', 'C'):
            src = db[f'{cc}{r}']; dst = db[f'{cc}{r-1}']
            dst.value = src.value
            if src.has_style:
                dst.font = copy.copy(src.font); dst.fill = copy.copy(src.fill)
                dst.number_format = src.number_format; dst.alignment = copy.copy(src.alignment)
                dst.border = copy.copy(src.border)
            src.value = None; src.fill = PatternFill(fill_type=None); src.border = Border()

    # ===================================================== new inputs (toggles)
    start = ip.max_row + 2
    S(ip, f'B{start}', 'Construction Interest & Refi', F_SECT, FILL_SECT, align=LEFT)
    for col in 'CDEF': ip[f'{col}{start}'].fill = FILL_SECT
    r = start + 1
    # RefiLagMonths
    S(ip, f'B{r}', 'Refi Lag Months', F_LBL, align=LEFT)
    S(ip, f'C{r}', 1, F_IN, FILL_IN, num=FMT_INT, align=RIGHT)
    S(ip, f'D{r}', 'months after stabilization/con-end to refinance (capped at maturity)', F_HINT, align=LEFT)
    repoint(wb, 'RefiLagMonths', f"'Inputs'!$C${r}"); r += 1
    # InterestMethod dropdown -> resolved current-pay flag in F
    S(ip, f'B{r}', 'Interest Method', F_LBL, align=LEFT)
    S(ip, f'C{r}', 'Current Pay', F_IN, FILL_IN, align=LEFT)
    dv = DataValidation(type='list', formula1='"Current Pay,Capitalized"', allow_blank=False, showErrorMessage=True)
    dv.promptTitle = 'Interest Method'; dv.prompt = 'Current Pay = pay interest from operations after conversion; Capitalized = fund all interest from reserve'
    ip.add_data_validation(dv); dv.add(ip[f'C{r}'])
    S(ip, f'D{r}', 'Capitalize pre-conversion; after the trigger, current-pay from operations', F_HINT, align=LEFT)
    S(ip, f'F{r}', f'=IF(C{r}="Current Pay",1,0)', F_LINK, num=FMT_INT, align=CENTER)
    repoint(wb, 'InterestMethodCurrentPay', f"'Inputs'!$F${r}"); r += 1
    # Conversion trigger month override (0 = stabilization)
    S(ip, f'B{r}', 'Interest Conversion Month', F_LBL, align=LEFT)
    S(ip, f'C{r}', 0, F_IN, FILL_IN, num=FMT_INT, align=RIGHT)
    S(ip, f'D{r}', '0 = convert at stabilization; >0 = manual month', F_HINT, align=LEFT)
    repoint(wb, 'IntConvMonthOverride', f"'Inputs'!$C${r}"); r += 1

    # DeliveryLagMonths 13 -> 19 (1C.3 reconcile)
    for rr in range(1, ip.max_row + 1):
        if ip.cell(rr, 2).value == 'DeliveryLagMonths':
            ip.cell(rr, 3).value = 19
            break

    # derived: IntConvMonth on Timeline
    tstart = tl.max_row + 2
    S(tl, f'B{tstart}', 'Interest Conversion Month', F_LBL, align=LEFT)
    S(tl, f'C{tstart}', '=IF(IntConvMonthOverride>0,IntConvMonthOverride,StabMonth)', F_FORM, num=FMT_INT, align=RIGHT)
    repoint(wb, 'IntConvMonth', f"'Timeline'!$C${tstart}")
    S(tl, f'B{tstart+1}', 'Construction End Month', F_LBL, align=LEFT)
    S(tl, f'C{tstart+1}', '=MS_Construction_End', F_FORM, num=FMT_INT, align=RIGHT)
    repoint(wb, 'ConEndMonth', f"'Timeline'!$C${tstart+1}")

    # ===================================================== 1C.4 refi timing
    tl['C24'].value = ('=IF(RefiFlag=1,MIN(ConMaturity,IF(RefiMonthOverride>0,RefiMonthOverride,'
                       'MAX(StabMonth,ConEndMonth)+RefiLagMonths)),0)')
    tl['C13'].value = '=IF(RefiFlag=1,RefiMonth,MS_Stabilization_End)'  # wire Recap milestone to real refi

    # ===================================================== 1C.7 untrended sizing
    cl['C6'].value = '=StabNOI_Untrended/ConMinDebtYield'
    cl['C7'].value = '=StabNOI_Untrended/(ConMinDSCR*ConRate)'
    cl['C16'].value = '=StabNOI_Untrended/ConCostCommit'
    cl['C17'].value = '=StabNOI_Untrended/(ConRate*ConCostCommit)'

    # ===================================================== 1C.8 dynamic index
    cl['C47'].value = '=INDEX(E32:NL32,ConPayoffMonth+1)'

    # ===================================================== 1C.6 LTC incl interest
    # reference capitalized-interest estimate (non-circular): balance per reference
    # debt = TargetLTC*ConCostBasis, drawn on the dev-spend curve, capitalized to the
    # conversion month.  Row 52 = reference balance; C54 = estimated reserve.
    # Row 52 = reference loan balance (debt = TargetLTC of the spend curve, drawn on
    # the dev-spend timing); row 53 = monthly capitalized interest on that balance,
    # accrued only while the loan is active and before conversion.  IntReserveEst =
    # Σ row 53 (always >= 0, non-circular — no reference to the sized commitment).
    for m, c in cols():
        p = get_column_letter(FIRST + m - 2) if m > 1 else None
        draw = f'TargetLTC*ConCostBasis*(DevSpend!{c}$34/DevSpend!$D$34)'
        capflag = f'IF(AND(Timeline!{c}$42=1,{c}$23<IntConvMonth),1,0)'
        if m == 1:
            cl[f'{c}53'].value = '=0'
            cl[f'{c}52'].value = f'={draw}'
        else:
            cl[f'{c}53'].value = f'={c}$30*{p}52*{capflag}'
            cl[f'{c}52'].value = f'={p}52+{draw}+{c}53'
        cl[f'{c}52'].number_format = FMT_MONEY
        cl[f'{c}53'].number_format = FMT_MONEY
    S(cl, 'B52', 'Ref Loan Balance (LTC interest estimate)', F_HINT, align=LEFT)
    S(cl, 'B53', 'Ref Capitalized Interest (monthly)', F_HINT, align=LEFT)
    S(cl, 'B54', 'Estimated Capitalized Interest (LTC sizing)', F_LBL, align=LEFT)
    S(cl, 'C54', '=SUM(E53:NL53)', F_FORM, num=FMT_MONEY, align=RIGHT)
    repoint(wb, 'IntReserveEst', "'ConLoan'!$C$54")
    # closed-form MaxByLTC incl. points, recourse, and capitalized interest:
    #   Commit <= (T*Base - (1-T)*R) / (1 - T*(points%+recourse%))
    cl['C5'].value = ('=(TargetLTC*ConCostBasis-(1-TargetLTC)*IntReserveEst)/'
                      '(1-TargetLTC*(ConPointsPct+RecourseFeePct))')

    # ===================================================== 1C.5 interest engine
    # capitalized interest (row 31): only while loan active AND pre-conversion (or
    # method = Capitalized).  Balance (row 32) unchanged structurally.
    for m, c in cols():
        p = get_column_letter(FIRST + m - 2) if m > 1 else None
        active = f'Timeline!{c}$42'
        capflag = f'IF(OR(InterestMethodCurrentPay=0,{c}$23<IntConvMonth),1,0)'
        if m == 1:
            cl[f'{c}31'].value = '=0'
        else:
            cl[f'{c}31'].value = f'=IF({active}=1,{c}$30*{p}32*{capflag},0)'
        # current-pay interest (row 50): while active, method=current pay, at/after trigger
        if m == 1:
            cl[f'{c}50'].value = '=0'
        else:
            cl[f'{c}50'].value = (f'=IF(AND({active}=1,InterestMethodCurrentPay=1,{c}$23>=IntConvMonth),'
                                  f'{c}$30*{p}32,0)')
        cl[f'{c}50'].number_format = FMT_MONEY
        # interest reserve balance drawdown (row 51)
        if m == 1:
            cl[f'{c}51'].value = '=InterestReserveFunded-E31'
        else:
            cl[f'{c}51'].value = f'={p}51-{c}31'
        cl[f'{c}51'].number_format = FMT_MONEY
    S(cl, 'B50', 'Interest — Current Pay (post-conversion)', F_HINT, align=LEFT)
    S(cl, 'B51', 'Interest Reserve Balance (drawdown)', F_HINT, align=LEFT)
    # totals + exhaustion
    S(cl, 'D50', '=SUM(E50:NL50)', F_FORM, num=FMT_MONEY, align=RIGHT)
    S(cl, 'B56', 'Current-Pay Interest — Total', F_LBL, align=LEFT)
    S(cl, 'C56', '=SUM(E50:NL50)', F_FORM, num=FMT_MONEY, align=RIGHT)
    repoint(wb, 'CurrentPayInterest', "'ConLoan'!$C$56")
    S(cl, 'B57', 'Min Interest Reserve Balance (>= 0)', F_LBL, align=LEFT)
    S(cl, 'C57', '=MIN(INDEX(E51:NL51,ConStart):INDEX(E51:NL51,ConPayoffMonth))', F_FORM, num=FMT_MONEY, align=RIGHT)
    repoint(wb, 'MinReserveBal', "'ConLoan'!$C$57")

    # levered CF (CashFlow row 6): subtract current-pay construction interest
    for m, c in cols():
        cur = cf[f'{c}6'].value
        if isinstance(cur, str) and 'ConLoan!' + c not in cur:
            cf[f'{c}6'].value = cur + f'-ConLoan!{c}$50'

    # ===================================================== diagnostics (hard)
    S(dg, 'B42', 'Engine Integrity (1C)', F_SECT, FILL_SECT, align=LEFT)
    for col in 'CD': dg[f'{col}42'].fill = FILL_SECT
    hard = [
        ('Final unit delivery ≤ construction end',
         '=(DeliverStart+CEILING(TotalUnits/UnitsDeliveredPerMo,1)-1)-ConEndMonth',
         '=IF((DeliverStart+CEILING(TotalUnits/UnitsDeliveredPerMo,1)-1)<=ConEndMonth,"PASS","FAIL")'),
        ('Stabilization ≥ construction end', '=StabMonth-ConEndMonth',
         '=IF(StabMonth>=ConEndMonth,"PASS","FAIL")'),
        ('Refi month ≥ construction end (if refi)', '=IF(RefiFlag=1,RefiMonth-ConEndMonth,0)',
         '=IF(RefiFlag=1,IF(RefiMonth>=ConEndMonth,"PASS","FAIL"),"N/A")'),
        ('Interest reserve not exhausted (≥ 0)', '=MinReserveBal',
         '=IF(MinReserveBal>=-1,"PASS","FAIL")'),
        ('Realized LTC ≤ target', '=LTC_Realized-TargetLTC',
         '=IF(LTC_Realized<=TargetLTC+0.005,"PASS","FAIL")'),
    ]
    r = 43
    for lab, cval, dform in hard:
        S(dg, f'B{r}', lab, F_LBL, align=LEFT)
        S(dg, f'C{r}', cval, F_FORM, num=FMT_MONEY, align=RIGHT)
        S(dg, f'D{r}', dform, F_LBLB, align=CENTER)
        r += 1
    hard_last = r - 1
    # rewire ErrorsFound to include the new hard checks
    err_row = None
    for rr in range(30, dg.max_row + 1):
        if dg.cell(rr, 2).value == 'TOTAL ERRORS FOUND':
            err_row = rr; break
    if err_row:
        dg.cell(err_row, 3).value = (f'=COUNTIF(D4:D25,"FAIL")+COUNTIF(D43:D{hard_last},"FAIL")'
                                     f'+COUNTIF(Proforma!$D$46:$D$46,"FAIL")')

    return {'ok': True, 'delivery_lag': 19}


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    print("applied:", build(wb))
    wb.save(dst)
    print("saved", dst)


if __name__ == '__main__':
    main()
