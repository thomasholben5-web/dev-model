"""Recalc an xlsx with LibreOffice (forces full recalc on load)."""
import sys, os, subprocess, time, openpyxl
from openpyxl.workbook.properties import CalcProperties
def recalc(src, outdir):
    wb=openpyxl.load_workbook(src)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    os.makedirs(outdir, exist_ok=True)
    tmp=os.path.join(outdir,"in.xlsx"); wb.save(tmp)
    out=os.path.join(outdir,"in.xlsx")
    # convert into a sibling dir to avoid overwriting the input mid-read
    od=os.path.join(outdir,"out"); os.makedirs(od, exist_ok=True)
    target=os.path.join(od,"in.xlsx")
    if os.path.exists(target): os.remove(target)
    env=dict(os.environ, HOME=outdir)
    p=subprocess.Popen(["soffice","--headless","--nolockcheck","--norestore",
        "--convert-to","xlsx:Calc MS Excel 2007 XML","--outdir",od, tmp],
        env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        stdin=subprocess.DEVNULL, start_new_session=True)
    for _ in range(120):
        if os.path.exists(target) and p.poll() is not None: break
        time.sleep(1)
    return target
if __name__=="__main__":
    src=sys.argv[1]; outdir=sys.argv[2]
    out=recalc(src,outdir)
    if not os.path.exists(out):
        print("RECALC FAILED - no output"); sys.exit(2)
    ERRS={'#REF!','#DIV/0!','#VALUE!','#N/A','#NAME?','#NUM!','#NULL!'}
    wb=openpyxl.load_workbook(out, data_only=True)
    bad=[f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets for row in ws.iter_rows()
         for c in row if isinstance(c.value,str) and any(e in c.value for e in ERRS)]
    print(f"ERRORS {len(bad)}")
    for b in bad[:40]: print("  ",b)
    names=sys.argv[3].split(",") if len(sys.argv)>3 else []
    wbf=openpyxl.load_workbook(src)
    for nm in names:
        dn=wbf.defined_names.get(nm)
        if not dn: print(f"  {nm}=<missing>"); continue
        for t,co in dn.destinations:
            print(f"  {nm} = {wb[t][co.replace('$','')].value}"); break
